"""ChemDB desktop client - refreshed Tkinter UI.

This file is a drop-in replacement for the original ``main.py``.  It keeps the
existing ``config.py`` and ``ssh_http.py`` backend contract unchanged.

Run the UI without a server for review::

    python main.py --demo
"""

from __future__ import annotations

import argparse
import base64
import binascii
import csv
import ctypes
import hashlib
import io
import json
import math
import os
import queue
import re
import sys
import threading
import time
import tkinter as tk
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from tkinter import colorchooser, filedialog, font as tkfont, messagebox, ttk
from typing import Any, Callable, Optional

from PIL import Image, ImageTk


IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")
FONT_FAMILY = "Microsoft YaHei UI"
MIN_FONT_SCALE = 0.9
MAX_FONT_SCALE = 1.4
DEFAULT_FONT_SCALE = 1.0
MIN_GRAPH_SPACING = 0.75
MAX_GRAPH_SPACING = 3.2
GRAPH_SPACING_STEP = 1.14
RICH_MARKER = "\n\n[CHEMDB_RICH_V1:"
RICH_END = "]"
STYLE_TAG_PREFIX = "chemstyle_"
KEYWORD_RULES_VERSION = 1
KEYWORD_TYPES = ("概念/物质", "关系词")
IMPORTANCE_LABELS = {1: "普通", 2: "重要", 3: "核心"}
SUPERSCRIPT_MAP = str.maketrans("0123456789+-=()", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾")
SUBSCRIPT_MAP = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
FONT_OPTIONS = ("Microsoft YaHei UI", "微软雅黑", "等线", "宋体", "黑体")
FONT_SIZE_OPTIONS = (9, 10, 11, 12, 14, 16, 18, 20, 24, 28, 32)
RELATION_TAXONOMY: dict[str, dict[str, tuple[str, ...]]] = {
    "基础反应": {
        "反应": ("反应", "发生反应", "化学反应"),
        "氧化还原": ("氧化还原", "氧化还原反应"),
        "氧化": ("氧化", "氧化反应"),
        "还原": ("还原", "还原反应"),
        "置换": ("置换", "置换反应"),
        "复分解": ("复分解", "复分解反应"),
        "中和": ("中和", "中和反应", "酸碱中和"),
        "水解": ("水解", "水解反应"),
        "电离": ("电离",),
        "解离": ("解离",),
        "歧化": ("歧化", "歧化反应"),
        "归中": ("归中", "归中反应"),
        "配位": ("配位", "配位反应"),
        "络合": ("络合", "络合反应"),
        "螯合": ("螯合", "螯合反应"),
        "沉淀": ("沉淀", "沉淀反应"),
        "共沉淀": ("共沉淀",),
        "溶解": ("溶解",),
        "生成": ("生成", "形成"),
        "分解": ("分解", "分解反应"),
        "转化": ("转化", "转变"),
    },
    "有机反应": {
        "取代": ("取代", "取代反应"),
        "加成": ("加成", "加成反应"),
        "消除": ("消除", "消除反应"),
        "酯化": ("酯化", "酯化反应"),
        "皂化": ("皂化", "皂化反应"),
        "缩合": ("缩合", "缩合反应"),
        "聚合": ("聚合", "聚合反应"),
        "裂解": ("裂解", "裂化"),
        "重排": ("重排", "重排反应"),
        "硝化": ("硝化", "硝化反应"),
        "磺化": ("磺化", "磺化反应"),
        "卤化": ("卤化", "卤化反应"),
        "氢化": ("氢化", "加氢"),
    },
    "催化与降解": {
        "催化": ("催化", "催化作用"),
        "光催化": ("光催化", "光催化作用"),
        "电催化": ("电催化", "电催化作用"),
        "生物催化": ("生物催化",),
        "催化降解": ("催化降解",),
        "光催化降解": ("光催化降解",),
        "生物降解": ("生物降解",),
        "光降解": ("光降解",),
        "氧化降解": ("氧化降解",),
        "还原降解": ("还原降解",),
        "降解": ("降解",),
        "矿化": ("矿化", "完全矿化"),
        "去除": ("去除", "脱除"),
    },
    "材料与界面": {
        "吸附": ("吸附", "表面吸附"),
        "脱附": ("脱附",),
        "离子交换": ("离子交换",),
        "包合": ("包合", "包合作用"),
        "嵌入": ("嵌入",),
        "嵌合": ("嵌合",),
        "插层": ("插层",),
        "负载": ("负载", "负载于"),
        "掺杂": ("掺杂",),
        "接枝": ("接枝",),
        "交联": ("交联",),
        "沉积": ("沉积",),
        "团聚": ("团聚", "聚集"),
        "分散": ("分散",),
        "溶胀": ("溶胀",),
        "复合": ("复合", "复合形成"),
        "氢键": ("氢键", "形成氢键"),
        "静电吸引": ("静电吸引", "静电作用"),
        "静电排斥": ("静电排斥",),
        "范德华作用": ("范德华作用", "范德华力"),
        "疏水作用": ("疏水作用", "疏水相互作用"),
        "π-π堆积": ("π-π堆积", "π–π堆积", "ππ堆积"),
        "相互作用": ("相互作用",),
    },
    "化工分离": {
        "吸收": ("吸收", "气体吸收"),
        "解吸": ("解吸", "气提"),
        "萃取": ("萃取", "液液萃取"),
        "浸取": ("浸取", "浸出"),
        "过滤": ("过滤",),
        "离心": ("离心", "离心分离"),
        "沉降": ("沉降",),
        "蒸发": ("蒸发",),
        "结晶": ("结晶", "重结晶"),
        "蒸馏": ("蒸馏",),
        "精馏": ("精馏",),
        "干燥": ("干燥",),
        "膜分离": ("膜分离",),
        "反渗透": ("反渗透",),
        "超滤": ("超滤",),
        "纳滤": ("纳滤",),
        "微滤": ("微滤",),
        "透析": ("透析",),
        "电渗析": ("电渗析",),
        "截留": ("截留",),
    },
    "传递与流动": {
        "传质": ("传质", "质量传递"),
        "传热": ("传热", "热量传递"),
        "导热": ("导热", "热传导"),
        "对流传热": ("对流传热",),
        "辐射传热": ("辐射传热", "热辐射"),
        "扩散": ("扩散", "分子扩散"),
        "对流": ("对流",),
        "渗透": ("渗透",),
        "透过": ("透过", "渗过"),
        "混合": ("混合",),
        "搅拌": ("搅拌",),
        "乳化": ("乳化",),
        "流动": ("流动",),
        "输送": ("输送",),
    },
    "影响关系": {
        "促进": ("促进", "有利于"),
        "抑制": ("抑制", "不利于"),
        "增强": ("增强",),
        "削弱": ("削弱",),
        "提高": ("提高", "提升", "增大", "增加"),
        "降低": ("降低", "减小", "减少"),
        "加快": ("加快", "加速"),
        "减慢": ("减慢", "减速"),
        "影响": ("影响",),
        "导致": ("导致", "引起", "造成"),
        "控制": ("控制",),
        "限制": ("限制",),
        "依赖": ("依赖", "取决于"),
        "正相关": ("正相关", "呈正相关"),
        "负相关": ("负相关", "呈负相关"),
    },
    "实验关系": {
        "对照": ("对照", "作为对照"),
        "重复": ("重复", "重复实验"),
        "验证": ("验证", "验证实验"),
        "比较": ("比较", "对比"),
        "优化": ("优化", "条件优化"),
        "改进": ("改进", "改善"),
        "延续": ("延续", "继续"),
        "衍生": ("衍生", "源自"),
        "复现": ("复现", "重现"),
        "前置": ("前置", "前置实验"),
        "后续": ("后续", "后续实验"),
    },
}

RELATION_LABEL_ALIASES: dict[str, str] = {}
RELATION_CATEGORY_BY_LABEL: dict[str, str] = {}
for _category, _relations in RELATION_TAXONOMY.items():
    for _canonical, _aliases in _relations.items():
        RELATION_CATEGORY_BY_LABEL[_canonical] = _category
        for _alias in (_canonical, *_aliases):
            RELATION_LABEL_ALIASES[_alias] = _canonical

RELATION_WORDS = tuple(sorted(RELATION_LABEL_ALIASES, key=len, reverse=True))


def _settings_path() -> Path:
    base = os.environ.get("APPDATA")
    if base:
        return Path(base) / "ChemDB" / "ui_settings.json"
    return Path.home() / ".chemdb" / "ui_settings.json"


def _keyword_rules_path() -> Path:
    return _settings_path().with_name("mindmap_keywords.json")


def _clamp_font_scale(value: Any) -> float:
    try:
        scale = float(value)
    except (TypeError, ValueError):
        scale = DEFAULT_FONT_SCALE
    return round(min(MAX_FONT_SCALE, max(MIN_FONT_SCALE, scale)), 2)


def load_font_scale() -> float:
    try:
        data = json.loads(_settings_path().read_text(encoding="utf-8"))
        return _clamp_font_scale(data.get("font_scale", DEFAULT_FONT_SCALE))
    except (OSError, ValueError, TypeError):
        return DEFAULT_FONT_SCALE


def save_font_scale(scale: float) -> None:
    try:
        path = _settings_path()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps({"font_scale": _clamp_font_scale(scale)}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except OSError:
        # Display preferences are optional and must never block the client.
        pass


def _validate_keyword_rule(value: Any) -> Optional[KeywordRule]:
    if not isinstance(value, dict):
        return None
    term = re.sub(r"\s+", " ", str(value.get("term", ""))).strip()
    label = re.sub(r"\s+", " ", str(value.get("label", ""))).strip()
    kind = str(value.get("kind", "concept"))
    if not term or len(term) > 48 or kind not in {"concept", "relation"}:
        return None
    try:
        importance = min(3, max(1, int(value.get("importance", 1))))
    except (TypeError, ValueError):
        importance = 1
    return KeywordRule(term, kind, label[:48], importance, bool(value.get("enabled", True)))


def load_keyword_rules() -> list[KeywordRule]:
    try:
        data = json.loads(_keyword_rules_path().read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return []
    if not isinstance(data, dict) or data.get("version") != KEYWORD_RULES_VERSION:
        return []
    rules: list[KeywordRule] = []
    seen: set[tuple[str, str]] = set()
    for item in data.get("keywords", []):
        rule = _validate_keyword_rule(item)
        if not rule:
            continue
        key = (rule.kind, rule.term.casefold())
        if key not in seen:
            seen.add(key)
            rules.append(rule)
    return rules


def save_keyword_rules(rules: list[KeywordRule]) -> None:
    path = _keyword_rules_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    payload = {
        "version": KEYWORD_RULES_VERSION,
        "keywords": [asdict(rule) for rule in rules],
    }
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def enable_high_dpi() -> None:
    """Opt into native Windows DPI rendering before the first Tk window exists."""
    if sys.platform != "win32":
        return
    try:
        ctypes.windll.user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4))
        return
    except (AttributeError, OSError):
        pass
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except (AttributeError, OSError):
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except (AttributeError, OSError):
            pass


enable_high_dpi()


@dataclass(frozen=True)
class Colors:
    canvas: str = "#F4F7FA"
    surface: str = "#FFFFFF"
    surface_alt: str = "#F8FAFC"
    border: str = "#DCE4EC"
    border_strong: str = "#C8D4DF"
    text: str = "#172B3A"
    text_muted: str = "#667A8A"
    text_soft: str = "#8B9AA6"
    primary: str = "#087F8C"
    primary_hover: str = "#076D78"
    primary_soft: str = "#E5F5F5"
    navy: str = "#123149"
    success: str = "#16835B"
    success_soft: str = "#E7F6EF"
    warning: str = "#B46A12"
    warning_soft: str = "#FFF4DF"
    danger: str = "#C43D4B"
    danger_soft: str = "#FDECEF"
    selection: str = "#DDF1F2"


C = Colors()


@dataclass(frozen=True)
class RichStyle:
    face: str = FONT_FAMILY
    size: int = 11
    bold: bool = False
    underline: bool = False
    foreground: str = C.text
    highlight: str = ""
    baseline: str = "normal"


@dataclass(frozen=True)
class KeywordRule:
    term: str
    kind: str = "concept"
    label: str = ""
    importance: int = 1
    enabled: bool = True

    @property
    def display_kind(self) -> str:
        return "关系词" if self.kind == "relation" else "概念/物质"

    @property
    def display_label(self) -> str:
        return self.label or self.term


@dataclass(frozen=True)
class Relation:
    source: str
    target: str
    label: str
    entry_id: int
    importance: int = 1


class DemoClient:
    """Small in-memory backend used only by ``--demo``."""

    def __init__(self) -> None:
        self._next_id = 4
        self._entries = [
            {
                "id": 1,
                "text": "亚甲基蓝与铅离子吸附实验\n记录不同 pH 条件下的吸光度变化。\n\n条件：25 ℃，反应 40 min。",
                "attachments": [
                    {"id": 101, "filename": "absorbance_data.xlsx", "size": 28416, "has_thumb": False},
                    {"id": 102, "filename": "reaction_sample.jpg", "size": 186320, "has_thumb": False},
                ],
            },
            {
                "id": 2,
                "text": "刚果红降解动力学\n比较不同催化剂投加量对降解速率的影响。",
                "attachments": [
                    {"id": 103, "filename": "kinetics.csv", "size": 8432, "has_thumb": False}
                ],
            },
            {
                "id": 3,
                "text": "罗丹明 B 基本性质\n分子式：C28H31ClN2O3\n保存条件：避光、干燥。",
                "attachments": [],
            },
        ]

    def connect(self) -> None:
        time.sleep(0.15)

    def close(self) -> None:
        return None

    def health(self) -> dict[str, str]:
        return {"status": "ok"}

    def list_entries(self, q: str = "", limit: int = 100, offset: int = 0) -> list[dict[str, Any]]:
        q = q.lower().strip()
        rows = [e for e in self._entries if not q or q in e["text"].lower()]
        return [self._copy(e) for e in rows[offset : offset + limit]]

    def get_entry(self, entry_id: int) -> dict[str, Any]:
        return self._copy(self._find(entry_id))

    def create_entry(self, text: str, file_paths: list[str]) -> dict[str, Any]:
        entry = {"id": self._next_id, "text": text, "attachments": []}
        self._next_id += 1
        self._entries.insert(0, entry)
        self._attach(entry, file_paths)
        return self._copy(entry)

    def update_entry(self, entry_id: int, text: str) -> dict[str, Any]:
        entry = self._find(entry_id)
        entry["text"] = text
        return self._copy(entry)

    def delete_entry(self, entry_id: int) -> dict[str, bool]:
        self._entries = [e for e in self._entries if e["id"] != entry_id]
        return {"ok": True}

    def add_attachments(self, entry_id: int, file_paths: list[str]) -> dict[str, Any]:
        entry = self._find(entry_id)
        self._attach(entry, file_paths)
        return self._copy(entry)

    def delete_attachment(self, att_id: int) -> dict[str, bool]:
        for entry in self._entries:
            entry["attachments"] = [a for a in entry["attachments"] if a["id"] != att_id]
        return {"ok": True}

    def download_attachment(self, att_id: int) -> bytes:
        if att_id == 101:
            try:
                from openpyxl import Workbook

                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "吸附关系"
                sheet.append(["物质A", "物质B", "关系", "pH"])
                sheet.append(["活性炭", "亚甲基蓝", "吸附", 7])
                output = io.BytesIO()
                workbook.save(output)
                return output.getvalue()
            except ImportError:
                return b""
        if att_id == 103:
            return "物质A,物质B,关系\n二氧化钛,刚果红,光催化降解\n".encode("utf-8-sig")
        return b"ChemDB demo attachment\n"

    def get_thumb(self, att_id: int) -> None:
        return None

    def _find(self, entry_id: int) -> dict[str, Any]:
        for entry in self._entries:
            if entry["id"] == entry_id:
                return entry
        raise KeyError(entry_id)

    @staticmethod
    def _copy(entry: dict[str, Any]) -> dict[str, Any]:
        return {**entry, "attachments": [dict(a) for a in entry.get("attachments", [])]}

    @staticmethod
    def _attach(entry: dict[str, Any], paths: list[str]) -> None:
        next_id = max([a["id"] for a in entry["attachments"]] + [200]) + 1
        for path in paths:
            try:
                size = os.path.getsize(path)
            except OSError:
                size = 0
            entry["attachments"].append(
                {
                    "id": next_id,
                    "filename": os.path.basename(path),
                    "size": size,
                    "has_thumb": False,
                }
            )
            next_id += 1


class ChemDBApp(tk.Tk):
    def __init__(self, demo: bool = False, font_scale: Optional[float] = None) -> None:
        super().__init__()
        self.demo = demo
        self.font_scale = _clamp_font_scale(load_font_scale() if font_scale is None else font_scale)
        self.client: Any = None
        self.entries: list[dict[str, Any]] = []
        self.all_entries: list[dict[str, Any]] = []
        self.current: Optional[dict[str, Any]] = None
        self._photo: Optional[ImageTk.PhotoImage] = None
        self._queue: queue.Queue[tuple[str, Any, Any]] = queue.Queue()
        self._search_job: Optional[str] = None
        self._busy_count = 0
        self._closing = False
        self._updating_editor = False
        self._dirty = False
        self._saved_text = ""
        self._rich_styles: dict[str, RichStyle] = {}
        self._view_zoom = 1.0
        self._relations: list[Relation] = []
        self._keyword_rules: list[KeywordRule] = load_keyword_rules()
        self._attachment_text_cache: dict[int, str] = {}
        self._attachment_parse_errors: dict[int, str] = {}
        self._graph_job: Optional[str] = None
        self._relation_analysis_job: Optional[str] = None
        self._graph_spacing = 1.0
        self._graph_pan_x = 0.0
        self._graph_pan_y = 0.0
        self._graph_drag_anchor: Optional[tuple[int, int]] = None
        self._restoring_selection = False

        self.search_var = tk.StringVar()
        self.title_var = tk.StringVar()
        self.editor_font_var = tk.StringVar(value=FONT_FAMILY)
        self.editor_size_var = tk.StringVar(value="11")
        self.view_zoom_var = tk.StringVar(value="100%")
        self.scale_var = tk.StringVar(value=f"{round(self.font_scale * 100)}%")
        self.connection_var = tk.StringVar(value="正在连接")
        self.status_var = tk.StringVar(value="正在初始化客户端…")
        self.count_var = tk.StringVar(value="0 条记录")
        self.selection_meta_var = tk.StringVar(value="尚未选择记录")
        self.preview_hint_var = tk.StringVar(value="选择附件后可在此预览")

        self._configure_window()
        self._configure_styles()
        self._build_ui()
        self._bind_shortcuts()
        self.title_var.trace_add("write", self._on_editor_changed)

        self.after(80, self._poll_queue)
        self.after(160, self._connect)
        self.after_idle(self._set_initial_sash)

    def _configure_window(self) -> None:
        self.title("ChemDB · 实验数据平台")
        dpi = float(self.winfo_fpixels("1i"))
        self.dpi_scale = min(2.5, max(1.0, dpi / 96.0)) if sys.platform == "win32" else 1.0
        self.layout_scale = 1.0 + (self.dpi_scale - 1.0) * 0.65
        if sys.platform == "win32":
            self.tk.call("tk", "scaling", dpi / 72.0)

        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()
        width = min(round(1280 * self.dpi_scale), max(1020, screen_w - 60))
        height = min(round(820 * self.dpi_scale), max(700, screen_h - 80))
        x = max(20, (screen_w - width) // 2)
        y = max(20, (screen_h - height) // 2)
        self.geometry(f"{width}x{height}+{x}+{y}")
        self.minsize(min(self._px(1080), width), min(self._px(700), height))
        self.configure(bg=C.canvas)
        self.option_add("*Font", self._font(11))
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def _font(self, size: float, weight: str = "normal") -> tuple[str, int, str]:
        return FONT_FAMILY, max(8, round(size * self.font_scale)), weight

    def _px(self, value: float) -> int:
        return max(1, round(value * self.layout_scale))

    def _padding(self, *values: int) -> tuple[int, ...]:
        return tuple(self._px(value) for value in values)

    def _row_px(self, value: float) -> int:
        return max(self._px(value), round(value * self.layout_scale * self.font_scale))

    def _configure_styles(self) -> None:
        style = ttk.Style(self)
        style.theme_use("clam")

        style.configure("App.TFrame", background=C.canvas)
        style.configure("Surface.TFrame", background=C.surface)
        style.configure("AltSurface.TFrame", background=C.surface_alt)
        style.configure("Header.TFrame", background=C.navy)

        style.configure("HeaderTitle.TLabel", background=C.navy, foreground="white", font=self._font(19, "bold"))
        style.configure("HeaderSub.TLabel", background=C.navy, foreground="#BFD1DC", font=self._font(10))
        style.configure("Title.TLabel", background=C.surface, foreground=C.text, font=self._font(15, "bold"))
        style.configure("Section.TLabel", background=C.surface, foreground=C.text, font=self._font(12, "bold"))
        style.configure("Body.TLabel", background=C.surface, foreground=C.text, font=self._font(11))
        style.configure("Muted.TLabel", background=C.surface, foreground=C.text_muted, font=self._font(10))
        style.configure("AltMuted.TLabel", background=C.surface_alt, foreground=C.text_muted, font=self._font(10))
        style.configure("Count.TLabel", background=C.surface, foreground=C.primary, font=self._font(10, "bold"))
        style.configure("Status.TLabel", background=C.canvas, foreground=C.text_muted, font=self._font(9.5))

        style.configure("TSeparator", background=C.border)
        style.configure("TEntry", fieldbackground=C.surface, foreground=C.text, bordercolor=C.border_strong, lightcolor=C.border_strong, darkcolor=C.border_strong, padding=self._padding(11, 9), font=self._font(11))
        style.map("TEntry", bordercolor=[("focus", C.primary)], lightcolor=[("focus", C.primary)], darkcolor=[("focus", C.primary)])

        style.configure("Primary.TButton", background=C.primary, foreground="white", borderwidth=0, focusthickness=0, padding=self._padding(16, 10), font=self._font(10, "bold"))
        style.map("Primary.TButton", background=[("active", C.primary_hover), ("disabled", "#A7C6C9")])
        style.configure("Secondary.TButton", background=C.surface, foreground=C.text, bordercolor=C.border_strong, lightcolor=C.border_strong, darkcolor=C.border_strong, padding=self._padding(13, 9), font=self._font(10))
        style.map("Secondary.TButton", background=[("active", C.surface_alt)])
        style.configure("Ghost.TButton", background=C.surface, foreground=C.primary, borderwidth=0, padding=self._padding(10, 8), font=self._font(10))
        style.map("Ghost.TButton", background=[("active", C.primary_soft)])
        style.configure("Danger.TButton", background=C.surface, foreground=C.danger, bordercolor="#F3C4CA", lightcolor="#F3C4CA", darkcolor="#F3C4CA", padding=self._padding(13, 9), font=self._font(10))
        style.map("Danger.TButton", background=[("active", C.danger_soft)])
        style.configure("Header.TButton", background=C.navy, foreground="#DDE9EF", bordercolor="#476276", lightcolor="#476276", darkcolor="#476276", padding=self._padding(13, 9), font=self._font(10))
        style.map("Header.TButton", background=[("active", "#1B405B")])
        style.configure("HeaderCompact.TButton", background=C.navy, foreground="#DDE9EF", bordercolor="#476276", lightcolor="#476276", darkcolor="#476276", padding=self._padding(8, 6), font=self._font(10, "bold"))
        style.map("HeaderCompact.TButton", background=[("active", "#1B405B")])

        style.configure("Records.Treeview", background=C.surface, fieldbackground=C.surface, foreground=C.text, borderwidth=0, rowheight=self._row_px(52), font=self._font(10.5))
        style.map("Records.Treeview", background=[("selected", C.selection)], foreground=[("selected", C.navy)])
        style.configure("Records.Treeview.Heading", background=C.surface_alt, foreground=C.text_muted, borderwidth=0, relief="flat", padding=self._padding(8, 9), font=self._font(9.5, "bold"))
        style.map("Records.Treeview.Heading", background=[("active", C.surface_alt)])

        style.configure("Attachments.Treeview", background=C.surface, fieldbackground=C.surface, foreground=C.text, borderwidth=0, rowheight=self._row_px(42), font=self._font(10))
        style.map("Attachments.Treeview", background=[("selected", C.primary_soft)], foreground=[("selected", C.navy)])
        style.configure("Attachments.Treeview.Heading", background=C.surface_alt, foreground=C.text_muted, borderwidth=0, relief="flat", padding=self._padding(8, 8), font=self._font(9.5, "bold"))

    def _build_ui(self) -> None:
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        self._build_header()

        workspace = ttk.Frame(self, style="App.TFrame", padding=(18, 18, 18, 10))
        workspace.grid(row=1, column=0, sticky="nsew")
        workspace.grid_rowconfigure(0, weight=1)
        workspace.grid_columnconfigure(0, weight=1)

        self.workspace_panes = tk.PanedWindow(
            workspace,
            orient=tk.HORIZONTAL,
            bg=C.canvas,
            borderwidth=0,
            sashwidth=self._px(9),
            sashrelief="flat",
            showhandle=True,
            opaqueresize=True,
        )
        self.workspace_panes.grid(row=0, column=0, sticky="nsew")
        records_host = tk.Frame(self.workspace_panes, bg=C.canvas)
        center_host = tk.Frame(self.workspace_panes, bg=C.canvas)
        map_host = tk.Frame(self.workspace_panes, bg=C.canvas)
        records_host.grid_rowconfigure(0, weight=1)
        records_host.grid_columnconfigure(0, weight=1)
        center_host.grid_rowconfigure(0, weight=1)
        center_host.grid_columnconfigure(0, weight=1)
        map_host.grid_rowconfigure(0, weight=1)
        map_host.grid_columnconfigure(0, weight=1)

        self.detail_panes = tk.PanedWindow(
            center_host,
            orient=tk.VERTICAL,
            bg=C.canvas,
            borderwidth=0,
            sashwidth=self._px(9),
            sashrelief="flat",
            showhandle=True,
            opaqueresize=True,
        )
        self.detail_panes.grid(row=0, column=0, sticky="nsew")
        editor_host = tk.Frame(self.detail_panes, bg=C.canvas)
        attachments_host = tk.Frame(self.detail_panes, bg=C.canvas)
        editor_host.grid_rowconfigure(0, weight=1)
        editor_host.grid_columnconfigure(0, weight=1)
        attachments_host.grid_rowconfigure(0, weight=1)
        attachments_host.grid_columnconfigure(0, weight=1)

        self._build_record_panel(records_host)
        self._build_detail_panel(editor_host)
        self._build_attachment_area(attachments_host)
        self._build_mindmap_panel(map_host)
        self.detail_panes.add(editor_host, minsize=self._px(330))
        self.detail_panes.add(attachments_host, minsize=self._px(190))
        self.workspace_panes.add(records_host, minsize=self._px(280), width=self._px(350))
        self.workspace_panes.add(center_host, minsize=self._px(580))
        self.workspace_panes.add(map_host, minsize=self._px(280), width=self._px(340))
        self._build_statusbar()
        self._show_empty_state()

    def _set_initial_sash(self) -> None:
        try:
            self.update_idletasks()
            total_width = self.workspace_panes.winfo_width()
            left_width = max(self._px(300), min(self._px(390), round(total_width * 0.23)))
            right_width = max(self._px(300), min(self._px(370), round(total_width * 0.22)))
            self.workspace_panes.sash_place(0, left_width, 0)
            self.workspace_panes.sash_place(1, total_width - right_width, 0)
            detail_height = self.detail_panes.winfo_height()
            self.detail_panes.sash_place(0, 0, round(detail_height * 0.67))
        except tk.TclError:
            pass

    def _build_header(self) -> None:
        header = ttk.Frame(self, style="Header.TFrame", padding=(22, 13))
        header.grid(row=0, column=0, sticky="ew")
        header.grid_columnconfigure(1, weight=1)

        self.brand_mark = tk.Label(
            header,
            text="C",
            width=2,
            height=1,
            bg="#1AA6A6",
            fg="white",
            font=self._font(19, "bold"),
        )
        self.brand_mark.grid(row=0, column=0, rowspan=2, sticky="w", padx=(0, 12))
        ttk.Label(header, text="ChemDB", style="HeaderTitle.TLabel").grid(row=0, column=1, sticky="sw")
        ttk.Label(header, text="团队实验数据与染料资料管理平台", style="HeaderSub.TLabel").grid(row=1, column=1, sticky="nw")

        ttk.Label(header, text="界面", style="HeaderSub.TLabel").grid(row=0, column=2, rowspan=2, padx=(8, 5))
        ttk.Button(header, text="A−", width=3, style="HeaderCompact.TButton", command=lambda: self._change_font_scale(-0.1)).grid(row=0, column=3, rowspan=2)
        self.scale_label = tk.Label(header, textvariable=self.scale_var, width=5, bg=C.navy, fg="#DDE9EF", font=self._font(9.5))
        self.scale_label.grid(row=0, column=4, rowspan=2, padx=4)
        ttk.Button(header, text="A+", width=3, style="HeaderCompact.TButton", command=lambda: self._change_font_scale(0.1)).grid(row=0, column=5, rowspan=2)

        self.connection_dot = tk.Label(header, text="●", bg=C.navy, fg="#F1B451", font=("Arial", round(12 * self.font_scale)))
        self.connection_dot.grid(row=0, column=6, rowspan=2, padx=(15, 5))
        self.connection_label = tk.Label(header, textvariable=self.connection_var, bg=C.navy, fg="#DDE9EF", font=self._font(10))
        self.connection_label.grid(row=0, column=7, rowspan=2, padx=(0, 14))
        ttk.Button(header, text="重新连接", style="Header.TButton", command=self._connect).grid(row=0, column=8, rowspan=2)

    def _build_record_panel(self, parent: ttk.Frame) -> None:
        panel = tk.Frame(parent, bg=C.surface, highlightbackground=C.border, highlightthickness=1)
        panel.grid(row=0, column=0, sticky="nsew", padx=(0, 7))
        panel.grid_columnconfigure(0, weight=1)
        panel.grid_rowconfigure(3, weight=1)

        head = ttk.Frame(panel, style="Surface.TFrame", padding=(16, 15, 16, 10))
        head.grid(row=0, column=0, sticky="ew")
        head.grid_columnconfigure(0, weight=1)
        ttk.Label(head, text="实验记录", style="Title.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(head, textvariable=self.count_var, style="Count.TLabel").grid(row=0, column=1, sticky="e")
        ttk.Label(head, text="按标题或正文快速检索团队资料", style="Muted.TLabel").grid(row=1, column=0, columnspan=2, sticky="w", pady=(4, 0))

        search = ttk.Frame(panel, style="Surface.TFrame", padding=(16, 4, 16, 10))
        search.grid(row=1, column=0, sticky="ew")
        search.grid_columnconfigure(0, weight=1)
        self.search_entry = ttk.Entry(search, textvariable=self.search_var)
        self.search_entry.grid(row=0, column=0, sticky="ew", padx=(0, 7))
        self.search_entry.insert(0, "")
        self.search_entry.bind("<KeyRelease>", self._schedule_search)
        self.search_entry.bind("<Return>", lambda _e: self.refresh())
        ttk.Button(search, text="搜索", style="Primary.TButton", command=self.refresh).grid(row=0, column=1)

        actions = ttk.Frame(panel, style="Surface.TFrame", padding=(16, 0, 16, 10))
        actions.grid(row=2, column=0, sticky="ew")
        actions.grid_columnconfigure(0, weight=1)
        ttk.Button(actions, text="＋ 新建实验记录", style="Primary.TButton", command=self._new_entry_dialog).grid(row=0, column=0, sticky="w")
        ttk.Button(actions, text="清除筛选", style="Ghost.TButton", command=self._show_all).grid(row=0, column=1, sticky="e")

        tree_host = ttk.Frame(panel, style="Surface.TFrame", padding=(8, 0, 8, 10))
        tree_host.grid(row=3, column=0, sticky="nsew")
        tree_host.grid_rowconfigure(0, weight=1)
        tree_host.grid_columnconfigure(0, weight=1)
        self.records = ttk.Treeview(tree_host, style="Records.Treeview", columns=("title", "files"), show="headings", selectmode="browse")
        self.records.heading("title", text="标题")
        self.records.heading("files", text="附件")
        self.records.column("title", width=235, minwidth=150, stretch=True, anchor="w")
        self.records.column("files", width=62, minwidth=58, stretch=False, anchor="center")
        self.records.grid(row=0, column=0, sticky="nsew")
        rscroll = ttk.Scrollbar(tree_host, orient="vertical", command=self.records.yview)
        rscroll.grid(row=0, column=1, sticky="ns")
        self.records.configure(yscrollcommand=rscroll.set)
        self.records.bind("<<TreeviewSelect>>", self._on_select)

    def _build_detail_panel(self, parent: ttk.Frame) -> None:
        self.detail_panel = tk.Frame(parent, bg=C.surface, highlightbackground=C.border, highlightthickness=1)
        self.detail_panel.grid(row=0, column=0, sticky="nsew")
        self.detail_panel.grid_columnconfigure(0, weight=1)
        self.detail_panel.grid_rowconfigure(2, weight=1)

        detail_head = ttk.Frame(self.detail_panel, style="Surface.TFrame", padding=(20, 15, 20, 12))
        detail_head.grid(row=0, column=0, sticky="ew")
        detail_head.grid_columnconfigure(0, weight=1)
        ttk.Label(detail_head, text="记录详情", style="Title.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(detail_head, textvariable=self.selection_meta_var, style="Muted.TLabel").grid(row=1, column=0, sticky="w", pady=(4, 0))
        self.save_button = ttk.Button(detail_head, text="保存修改  Ctrl+S", style="Primary.TButton", command=self._save_text)
        self.save_button.grid(row=0, column=1, rowspan=2, padx=(10, 0))
        self.delete_entry_button = ttk.Button(detail_head, text="删除记录", style="Danger.TButton", command=self._delete_entry)
        self.delete_entry_button.grid(row=0, column=2, rowspan=2, padx=(8, 0))

        title_box = ttk.Frame(self.detail_panel, style="Surface.TFrame", padding=(20, 10, 20, 8))
        title_box.grid(row=1, column=0, sticky="ew")
        title_box.grid_columnconfigure(0, weight=1)
        ttk.Label(title_box, text="实验标题", style="Section.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 7))
        self.title_entry = ttk.Entry(title_box, textvariable=self.title_var, font=self._font(12, "bold"))
        self.title_entry.grid(row=1, column=0, sticky="ew")

        body_box = ttk.Frame(self.detail_panel, style="Surface.TFrame", padding=(20, 6, 20, 14))
        body_box.grid(row=2, column=0, sticky="nsew")
        body_box.grid_rowconfigure(2, weight=1)
        body_box.grid_columnconfigure(0, weight=1)
        ttk.Label(body_box, text="实验内容", style="Section.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 7))
        self._build_rich_text_toolbar(body_box)
        text_host = tk.Frame(body_box, bg=C.surface, highlightbackground=C.border_strong, highlightthickness=1)
        text_host.grid(row=2, column=0, sticky="nsew")
        text_host.grid_rowconfigure(0, weight=1)
        text_host.grid_columnconfigure(0, weight=1)
        self.body_text = tk.Text(
            text_host,
            wrap="word",
            undo=True,
            exportselection=False,
            relief="flat",
            borderwidth=0,
            bg=C.surface,
            fg=C.text,
            insertbackground=C.primary,
            selectbackground=C.selection,
            font=self._font(11),
            padx=12,
            pady=10,
            spacing1=2,
            spacing3=3,
        )
        self.body_text.grid(row=0, column=0, sticky="nsew")
        self.body_text.bind("<<Modified>>", self._on_editor_text_modified)
        bscroll = ttk.Scrollbar(text_host, orient="vertical", command=self.body_text.yview)
        bscroll.grid(row=0, column=1, sticky="ns")
        self.body_text.configure(yscrollcommand=bscroll.set)

        self.empty_state = ttk.Frame(self.detail_panel, style="Surface.TFrame")
        self.empty_state.place(relx=0, rely=0, relwidth=1, relheight=1)
        center = ttk.Frame(self.empty_state, style="Surface.TFrame")
        center.place(relx=0.5, rely=0.44, anchor="center")
        self.empty_icon = tk.Label(center, text="⌁", bg=C.surface, fg=C.primary, font=("Arial", round(36 * self.font_scale)))
        self.empty_icon.pack(pady=(0, 8))
        ttk.Label(center, text="选择一条实验记录", style="Title.TLabel").pack()
        ttk.Label(center, text="从左侧打开记录，或新建一条实验记录", style="Muted.TLabel").pack(pady=(6, 16))
        ttk.Button(center, text="＋ 新建记录", style="Primary.TButton", command=self._new_entry_dialog).pack()

    def _build_rich_text_toolbar(self, parent: ttk.Frame) -> None:
        toolbar = ttk.Frame(parent, style="AltSurface.TFrame", padding=(7, 6))
        toolbar.grid(row=1, column=0, sticky="ew", pady=(0, 7))
        toolbar.grid_columnconfigure(10, weight=1)

        installed_fonts = sorted(set(tkfont.families(self)), key=str.casefold)
        preferred_fonts = [name for name in FONT_OPTIONS if name in installed_fonts]
        font_choices = tuple(preferred_fonts + [name for name in installed_fonts if name not in preferred_fonts]) or FONT_OPTIONS

        self.font_combo = ttk.Combobox(
            toolbar,
            textvariable=self.editor_font_var,
            values=font_choices,
            state="readonly",
            width=15,
            font=self._font(9.5),
        )
        self.font_combo.grid(row=0, column=0, padx=(0, 5))
        self.font_combo.bind("<<ComboboxSelected>>", lambda _e: self._apply_style_change(face=self.editor_font_var.get()))

        self.size_combo = ttk.Combobox(
            toolbar,
            textvariable=self.editor_size_var,
            values=tuple(str(value) for value in FONT_SIZE_OPTIONS),
            state="readonly",
            width=4,
            font=self._font(9.5),
        )
        self.size_combo.grid(row=0, column=1, padx=(0, 5))
        self.size_combo.bind("<<ComboboxSelected>>", lambda _e: self._apply_style_change(size=int(self.editor_size_var.get())))

        ttk.Button(toolbar, text="B", width=3, style="Secondary.TButton", command=lambda: self._toggle_style("bold")).grid(row=0, column=2, padx=2)
        ttk.Button(toolbar, text="U̲", width=3, style="Secondary.TButton", command=lambda: self._toggle_style("underline")).grid(row=0, column=3, padx=2)
        ttk.Button(toolbar, text="x²", width=3, style="Secondary.TButton", command=lambda: self._toggle_baseline("super")).grid(row=0, column=4, padx=2)
        ttk.Button(toolbar, text="x₂", width=3, style="Secondary.TButton", command=lambda: self._toggle_baseline("sub")).grid(row=0, column=5, padx=2)
        ttk.Button(toolbar, text="文字色", style="Secondary.TButton", command=self._pick_text_color).grid(row=0, column=6, padx=2)
        ttk.Button(toolbar, text="荧光笔", style="Secondary.TButton", command=self._pick_highlight).grid(row=0, column=7, padx=2)
        ttk.Button(toolbar, text="清除格式", style="Ghost.TButton", command=self._clear_selection_style).grid(row=0, column=8, padx=(2, 7))

        view_tools = ttk.Frame(toolbar, style="AltSurface.TFrame")
        view_tools.grid(row=0, column=11, sticky="e")
        ttk.Label(view_tools, text="阅读", style="AltMuted.TLabel").pack(side="left", padx=(0, 3))
        ttk.Button(view_tools, text="−", width=3, style="Ghost.TButton", command=lambda: self._change_view_zoom(-0.1)).pack(side="left")
        ttk.Label(view_tools, textvariable=self.view_zoom_var, style="AltMuted.TLabel", width=5, anchor="center").pack(side="left")
        ttk.Button(view_tools, text="＋", width=3, style="Ghost.TButton", command=lambda: self._change_view_zoom(0.1)).pack(side="left")

    def _build_attachment_area(self, parent: tk.Widget) -> None:
        self.attachment_panel = tk.Frame(parent, bg=C.surface, highlightbackground=C.border, highlightthickness=1)
        self.attachment_panel.grid(row=0, column=0, sticky="nsew")
        self.attachment_panel.grid_rowconfigure(0, weight=1)
        self.attachment_panel.grid_columnconfigure(0, weight=1)
        area = ttk.Frame(self.attachment_panel, style="Surface.TFrame", padding=(20, 12, 20, 16))
        area.grid(row=0, column=0, sticky="nsew")
        area.grid_columnconfigure(0, weight=3)
        area.grid_columnconfigure(1, weight=2)
        area.grid_rowconfigure(1, weight=1)

        ttk.Label(area, text="附件与原始数据", style="Section.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 8))
        tools = ttk.Frame(area, style="Surface.TFrame")
        tools.grid(row=0, column=1, sticky="e", pady=(0, 8))
        ttk.Button(tools, text="添加", style="Ghost.TButton", command=self._add_attachments).pack(side="left")
        ttk.Button(tools, text="下载", style="Ghost.TButton", command=self._download).pack(side="left")
        ttk.Button(tools, text="删除", style="Danger.TButton", command=self._delete_attachment).pack(side="left", padx=(5, 0))

        table_host = tk.Frame(area, bg=C.surface, highlightbackground=C.border, highlightthickness=1)
        table_host.grid(row=1, column=0, sticky="nsew", padx=(0, 10))
        table_host.grid_rowconfigure(0, weight=1)
        table_host.grid_columnconfigure(0, weight=1)
        self.attachments = ttk.Treeview(table_host, style="Attachments.Treeview", columns=("name", "type", "size"), show="headings", selectmode="browse", height=4)
        self.attachments.heading("name", text="文件名")
        self.attachments.heading("type", text="类型")
        self.attachments.heading("size", text="大小")
        self.attachments.column("name", minwidth=170, width=280, stretch=True)
        self.attachments.column("type", width=72, stretch=False, anchor="center")
        self.attachments.column("size", width=78, stretch=False, anchor="e")
        self.attachments.grid(row=0, column=0, sticky="nsew")
        ascroll = ttk.Scrollbar(table_host, orient="vertical", command=self.attachments.yview)
        ascroll.grid(row=0, column=1, sticky="ns")
        self.attachments.configure(yscrollcommand=ascroll.set)
        self.attachments.bind("<<TreeviewSelect>>", lambda _e: self._preview())
        self.attachments.bind("<Double-Button-1>", lambda _e: self._preview())

        self.preview_box = tk.Frame(area, bg=C.surface_alt, highlightbackground=C.border, highlightthickness=1)
        self.preview_box.grid(row=1, column=1, sticky="nsew")
        self.preview_box.grid_rowconfigure(0, weight=1)
        self.preview_box.grid_columnconfigure(0, weight=1)
        self.preview_label = tk.Label(
            self.preview_box,
            textvariable=self.preview_hint_var,
            bg=C.surface_alt,
            fg=C.text_muted,
            font=self._font(10),
            justify="center",
            wraplength=220,
        )
        self.preview_label.grid(row=0, column=0, sticky="nsew", padx=10, pady=10)

    def _build_mindmap_panel(self, parent: tk.Widget) -> None:
        self.mindmap_panel = tk.Frame(parent, bg=C.surface, highlightbackground=C.border, highlightthickness=1)
        self.mindmap_panel.grid(row=0, column=0, sticky="nsew")
        self.mindmap_panel.grid_rowconfigure(1, weight=1)
        self.mindmap_panel.grid_columnconfigure(0, weight=1)

        head = ttk.Frame(self.mindmap_panel, style="Surface.TFrame", padding=(16, 15, 16, 11))
        head.grid(row=0, column=0, sticky="ew")
        head.grid_columnconfigure(0, weight=1)
        ttk.Label(head, text="关系思维导图", style="Title.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Button(head, text="关键词", style="Ghost.TButton", command=self._open_keyword_manager).grid(row=0, column=1, sticky="e")
        ttk.Button(head, text="重置视图", style="Ghost.TButton", command=self._reset_graph_view).grid(row=0, column=2, sticky="e", padx=(4, 0))
        ttk.Button(head, text="重新分析", style="Ghost.TButton", command=self._rebuild_relation_graph).grid(row=0, column=3, sticky="e", padx=(4, 0))
        ttk.Label(head, text="自动提取物质、条件与实验之间的关系", style="Muted.TLabel").grid(row=1, column=0, columnspan=4, sticky="w", pady=(4, 0))

        self.mindmap_canvas = tk.Canvas(
            self.mindmap_panel,
            bg=C.surface,
            highlightthickness=0,
            borderwidth=0,
        )
        self.mindmap_canvas.grid(row=1, column=0, sticky="nsew", padx=8)
        self.mindmap_canvas.bind("<Configure>", self._schedule_graph_draw)
        self.mindmap_canvas.bind("<MouseWheel>", self._on_graph_mousewheel)
        self.mindmap_canvas.bind("<Button-4>", self._on_graph_mousewheel)
        self.mindmap_canvas.bind("<Button-5>", self._on_graph_mousewheel)
        self.mindmap_canvas.bind("<ButtonPress-1>", self._on_graph_drag_start)
        self.mindmap_canvas.bind("<B1-Motion>", self._on_graph_drag_move)
        self.mindmap_canvas.bind("<ButtonRelease-1>", self._on_graph_drag_end)

        self.graph_status_var = tk.StringVar(value="选择记录后自动生成")
        footer = ttk.Frame(self.mindmap_panel, style="Surface.TFrame", padding=(14, 8, 14, 12))
        footer.grid(row=2, column=0, sticky="ew")
        ttk.Label(footer, textvariable=self.graph_status_var, style="Muted.TLabel").pack(anchor="w")
        ttk.Label(footer, text="滚轮调整间距 · 按住鼠标左键拖动画布", style="Muted.TLabel").pack(anchor="w", pady=(3, 0))

    @staticmethod
    def _normalize_chemical_notation(value: str) -> str:
        """Accept Fe^{3+}/SO_{4}^{2-} as friendly input and store Unicode."""
        text = value.strip()
        text = re.sub(r"_\{([^{}]+)\}", lambda match: match.group(1).translate(SUBSCRIPT_MAP), text)
        text = re.sub(r"\^\{([^{}]+)\}", lambda match: match.group(1).translate(SUPERSCRIPT_MAP), text)
        return text

    def _open_keyword_manager(self) -> None:
        dialog = tk.Toplevel(self)
        dialog.title("思维导图关键词")
        dialog.geometry(f"{self._px(820)}x{self._px(610)}")
        dialog.minsize(self._px(690), self._px(520))
        dialog.transient(self)
        dialog.grab_set()
        dialog.configure(bg=C.canvas)

        root = ttk.Frame(dialog, style="App.TFrame", padding=18)
        root.pack(fill="both", expand=True)
        root.grid_columnconfigure(0, weight=1)
        root.grid_rowconfigure(1, weight=1)
        ttk.Label(root, text="自定义识别关键词", style="Title.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(root, text="可添加物质/概念或新的关系词；重要性用于中心选择与导图强调。", style="Muted.TLabel").grid(row=0, column=1, sticky="e")

        table_host = tk.Frame(root, bg=C.surface, highlightbackground=C.border, highlightthickness=1)
        table_host.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(12, 12))
        table_host.grid_rowconfigure(0, weight=1)
        table_host.grid_columnconfigure(0, weight=1)
        tree = ttk.Treeview(table_host, columns=("kind", "term", "label", "importance"), show="headings", selectmode="browse")
        for column, title, width in (("kind", "类型", 100), ("term", "识别文字", 220), ("label", "导图显示/规范关系", 230), ("importance", "重要性", 90)):
            tree.heading(column, text=title)
            tree.column(column, width=self._px(width), minwidth=self._px(70), stretch=column in {"term", "label"})
        tree.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(table_host, orient="vertical", command=tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        tree.configure(yscrollcommand=scrollbar.set)

        form = ttk.Frame(root, style="Surface.TFrame", padding=12)
        form.grid(row=2, column=0, columnspan=2, sticky="ew")
        form.grid_columnconfigure(1, weight=1)
        form.grid_columnconfigure(3, weight=1)
        kind_var = tk.StringVar(value=KEYWORD_TYPES[0])
        term_var = tk.StringVar()
        label_var = tk.StringVar()
        importance_var = tk.StringVar(value="普通")
        editing_index: list[Optional[int]] = [None]

        ttk.Label(form, text="类型", style="Muted.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 6))
        kind_combo = ttk.Combobox(form, textvariable=kind_var, values=KEYWORD_TYPES, state="readonly", width=11)
        kind_combo.grid(row=0, column=1, sticky="w", padx=(0, 14))
        ttk.Label(form, text="重要性", style="Muted.TLabel").grid(row=0, column=2, sticky="w", padx=(0, 6))
        ttk.Combobox(form, textvariable=importance_var, values=tuple(IMPORTANCE_LABELS.values()), state="readonly", width=9).grid(row=0, column=3, sticky="w")
        ttk.Label(form, text="识别文字", style="Muted.TLabel").grid(row=1, column=0, sticky="w", pady=(10, 0), padx=(0, 6))
        term_entry = ttk.Entry(form, textvariable=term_var)
        term_entry.grid(row=1, column=1, sticky="ew", pady=(10, 0), padx=(0, 14))
        ttk.Label(form, text="显示名称", style="Muted.TLabel").grid(row=1, column=2, sticky="w", pady=(10, 0), padx=(0, 6))
        ttk.Entry(form, textvariable=label_var).grid(row=1, column=3, sticky="ew", pady=(10, 0))
        ttk.Label(form, text="化学式可输入 Fe^{3+}、SO_{4}^{2-}，保存后自动转换为 Fe³⁺、SO₄²⁻。关系词的显示名称用于归一化标签。", style="Muted.TLabel").grid(row=2, column=0, columnspan=4, sticky="w", pady=(9, 0))

        def refresh_tree(select: Optional[int] = None) -> None:
            tree.delete(*tree.get_children())
            for index, rule in enumerate(self._keyword_rules):
                tree.insert("", "end", iid=str(index), values=(rule.display_kind, rule.term, rule.display_label, IMPORTANCE_LABELS[rule.importance]))
            if select is not None and tree.exists(str(select)):
                tree.selection_set(str(select))
                tree.focus(str(select))

        def clear_form() -> None:
            editing_index[0] = None
            kind_var.set(KEYWORD_TYPES[0])
            term_var.set("")
            label_var.set("")
            importance_var.set("普通")
            term_entry.focus_set()

        def load_selected(_event: Optional[tk.Event] = None) -> None:
            selection = tree.selection()
            if not selection:
                return
            index = int(selection[0])
            rule = self._keyword_rules[index]
            editing_index[0] = index
            kind_var.set(rule.display_kind)
            term_var.set(rule.term)
            label_var.set(rule.label)
            importance_var.set(IMPORTANCE_LABELS[rule.importance])

        def save_rule() -> None:
            term = self._normalize_chemical_notation(term_var.get())
            label = self._normalize_chemical_notation(label_var.get())
            if not term:
                messagebox.showwarning("缺少关键词", "请输入需要识别的文字。", parent=dialog)
                return
            kind = "relation" if kind_var.get() == "关系词" else "concept"
            importance = next((value for value, name in IMPORTANCE_LABELS.items() if name == importance_var.get()), 1)
            duplicate = next((i for i, item in enumerate(self._keyword_rules) if item.kind == kind and item.term.casefold() == term.casefold() and i != editing_index[0]), None)
            if duplicate is not None:
                messagebox.showwarning("关键词已存在", "同类型下已经存在这个关键词。", parent=dialog)
                return
            rule = KeywordRule(term=term, kind=kind, label=label, importance=importance)
            if editing_index[0] is None:
                self._keyword_rules.append(rule)
                selected = len(self._keyword_rules) - 1
            else:
                selected = editing_index[0]
                self._keyword_rules[selected] = rule
            try:
                save_keyword_rules(self._keyword_rules)
            except OSError as exc:
                messagebox.showerror("保存失败", str(exc), parent=dialog)
                return
            refresh_tree(selected)
            editing_index[0] = selected
            self._rebuild_relation_graph()
            self.status_var.set(f"已保存关键词“{term}”")

        def delete_rule() -> None:
            selection = tree.selection()
            if not selection:
                return
            index = int(selection[0])
            rule = self._keyword_rules[index]
            if not messagebox.askyesno("删除关键词", f"确定删除“{rule.term}”吗？", parent=dialog):
                return
            del self._keyword_rules[index]
            try:
                save_keyword_rules(self._keyword_rules)
            except OSError as exc:
                messagebox.showerror("保存失败", str(exc), parent=dialog)
                return
            clear_form()
            refresh_tree()
            self._rebuild_relation_graph()

        tree.bind("<<TreeviewSelect>>", load_selected)
        buttons = ttk.Frame(root, style="App.TFrame")
        buttons.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        ttk.Button(buttons, text="新建", style="Secondary.TButton", command=clear_form).pack(side="left")
        ttk.Button(buttons, text="删除", style="Danger.TButton", command=delete_rule).pack(side="left", padx=(8, 0))
        ttk.Button(buttons, text="保存关键词", style="Primary.TButton", command=save_rule).pack(side="right")
        ttk.Button(buttons, text="关闭", style="Secondary.TButton", command=dialog.destroy).pack(side="right", padx=(0, 8))
        refresh_tree()
        term_entry.focus_set()

    def _build_statusbar(self) -> None:
        status = ttk.Frame(self, style="App.TFrame", padding=(20, 0, 20, 10))
        status.grid(row=2, column=0, sticky="ew")
        status.grid_columnconfigure(0, weight=1)
        ttk.Label(status, textvariable=self.status_var, style="Status.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(status, text="Ctrl+N 新建   Ctrl+S 保存   Ctrl+F 搜索   F5 刷新", style="Status.TLabel").grid(row=0, column=1, sticky="e")

    def _bind_shortcuts(self) -> None:
        self.bind_all("<Control-n>", lambda _e: self._new_entry_dialog())
        self.bind_all("<Control-s>", lambda _e: self._save_text())
        self.bind_all("<Control-f>", lambda _e: self._focus_search())
        self.bind_all("<F5>", lambda _e: self.refresh())
        self.bind_all("<Control-plus>", lambda _e: self._change_font_scale(0.1))
        self.bind_all("<Control-equal>", lambda _e: self._change_font_scale(0.1))
        self.bind_all("<Control-minus>", lambda _e: self._change_font_scale(-0.1))
        self.bind_all("<Control-0>", lambda _e: self._set_font_scale(DEFAULT_FONT_SCALE))

    def _change_font_scale(self, delta: float) -> str:
        return self._set_font_scale(self.font_scale + delta)

    def _set_font_scale(self, scale: float) -> str:
        new_scale = _clamp_font_scale(scale)
        if new_scale == self.font_scale:
            return "break"
        self.font_scale = new_scale
        self.scale_var.set(f"{round(new_scale * 100)}%")
        self._configure_styles()
        self.option_add("*Font", self._font(11))
        self.brand_mark.configure(font=self._font(19, "bold"))
        self.scale_label.configure(font=self._font(9.5))
        self.connection_dot.configure(font=("Arial", round(12 * self.font_scale)))
        self.connection_label.configure(font=self._font(10))
        self.title_entry.configure(font=self._font(12, "bold"))
        self.font_combo.configure(font=self._font(9.5))
        self.size_combo.configure(font=self._font(9.5))
        self.body_text.configure(font=self._font(11 * self._view_zoom))
        for tag, style in self._rich_styles.items():
            self._configure_style_tag(tag, style)
        self.preview_label.configure(font=self._font(10))
        self.empty_icon.configure(font=("Arial", round(36 * self.font_scale)))
        save_font_scale(new_scale)
        self.status_var.set(f"界面字号已调整为 {round(new_scale * 100)}%")
        return "break"

    def _on_editor_text_modified(self, _event: tk.Event) -> None:
        if self.body_text.edit_modified():
            self.body_text.edit_modified(False)
            self._on_editor_changed()

    def _on_editor_changed(self, *_args: Any, force: bool = False) -> None:
        if self._updating_editor or not self.current:
            return
        self._schedule_relation_analysis()
        current_text = self._compose_current_text()
        dirty = current_text != self._saved_text
        if dirty == self._dirty and not force:
            return
        self._dirty = dirty
        self._update_selection_meta()
        self.save_button.configure(text="保存修改 *  Ctrl+S" if dirty else "保存修改  Ctrl+S")

    def _update_selection_meta(self) -> None:
        if not self.current:
            self.selection_meta_var.set("尚未选择记录")
            return
        attachments = self.current.get("attachments", []) or []
        suffix = "   ·   有未保存修改" if self._dirty else ""
        self.selection_meta_var.set(f"记录 #{self.current.get('id', '—')}   ·   {len(attachments)} 个附件{suffix}")

    def _selected_text_range(self) -> Optional[tuple[str, str]]:
        try:
            return self.body_text.index("sel.first"), self.body_text.index("sel.last")
        except tk.TclError:
            self.status_var.set("请先在实验内容中选中需要设置格式的文字")
            self.body_text.focus_set()
            return None

    def _style_at(self, index: str) -> RichStyle:
        for tag in reversed(self.body_text.tag_names(index)):
            if tag.startswith(STYLE_TAG_PREFIX) and tag in self._rich_styles:
                return self._rich_styles[tag]
        return RichStyle()

    def _style_tag(self, style: RichStyle) -> str:
        encoded = json.dumps(asdict(style), ensure_ascii=False, sort_keys=True).encode("utf-8")
        tag = STYLE_TAG_PREFIX + hashlib.sha1(encoded).hexdigest()[:14]
        self._rich_styles[tag] = style
        self._configure_style_tag(tag, style)
        return tag

    def _configure_style_tag(self, tag: str, style: RichStyle) -> None:
        baseline = style.baseline if style.baseline in {"normal", "super", "sub"} else "normal"
        factor = 0.78 if baseline != "normal" else 1.0
        size = max(7, round(style.size * self.font_scale * self._view_zoom * factor))
        offset = 0
        if baseline == "super":
            offset = max(2, round(style.size * self.font_scale * self._view_zoom * 0.30))
        elif baseline == "sub":
            offset = -max(2, round(style.size * self.font_scale * self._view_zoom * 0.18))
        self.body_text.tag_configure(
            tag,
            font=(style.face, size, "bold" if style.bold else "normal"),
            underline=1 if style.underline else 0,
            foreground=style.foreground,
            background=style.highlight or C.surface,
            offset=offset,
        )
        self.body_text.tag_raise("sel")

    def _selection_style_runs(self, start: str, end: str) -> list[tuple[str, str, RichStyle]]:
        length = int(self.body_text.count(start, end, "chars")[0])
        if length <= 0:
            return []
        runs: list[tuple[str, str, RichStyle]] = []
        run_start = 0
        current = self._style_at(start)
        for offset in range(1, length):
            index = self.body_text.index(f"{start}+{offset}c")
            style = self._style_at(index)
            if style != current:
                runs.append((self.body_text.index(f"{start}+{run_start}c"), index, current))
                run_start = offset
                current = style
        runs.append((self.body_text.index(f"{start}+{run_start}c"), end, current))
        return runs

    def _apply_style_change(self, **changes: Any) -> None:
        selected = self._selected_text_range()
        if not selected:
            return
        start, end = selected
        runs = self._selection_style_runs(start, end)
        for tag in tuple(self._rich_styles):
            self.body_text.tag_remove(tag, start, end)
        for run_start, run_end, style in runs:
            updated = replace(style, **changes)
            if updated != RichStyle():
                self.body_text.tag_add(self._style_tag(updated), run_start, run_end)
        self.body_text.tag_add("sel", start, end)
        self._mark_rich_change()

    def _toggle_style(self, attribute: str) -> None:
        selected = self._selected_text_range()
        if not selected:
            return
        current = self._style_at(selected[0])
        self._apply_style_change(**{attribute: not bool(getattr(current, attribute))})

    def _toggle_baseline(self, baseline: str) -> None:
        selected = self._selected_text_range()
        if not selected:
            return
        current = self._style_at(selected[0])
        self._apply_style_change(baseline="normal" if current.baseline == baseline else baseline)

    def _pick_text_color(self) -> None:
        selected = self._selected_text_range()
        if not selected:
            return
        current = self._style_at(selected[0])
        _rgb, color = colorchooser.askcolor(current.foreground, title="选择文字颜色", parent=self)
        if color:
            self._apply_style_change(foreground=color.upper())

    def _pick_highlight(self) -> None:
        if not self._selected_text_range():
            return
        _rgb, color = colorchooser.askcolor("#FFF2A8", title="选择荧光标记颜色", parent=self)
        if color:
            self._apply_style_change(highlight=color.upper())

    def _clear_selection_style(self) -> None:
        selected = self._selected_text_range()
        if not selected:
            return
        start, end = selected
        for tag in tuple(self._rich_styles):
            self.body_text.tag_remove(tag, start, end)
        self.body_text.tag_add("sel", start, end)
        self._mark_rich_change()

    def _mark_rich_change(self) -> None:
        self._on_editor_changed(force=True)
        self.body_text.focus_set()

    def _change_view_zoom(self, delta: float) -> None:
        self._view_zoom = round(min(1.8, max(0.8, self._view_zoom + delta)), 1)
        self.view_zoom_var.set(f"{round(self._view_zoom * 100)}%")
        self.body_text.configure(font=self._font(11 * self._view_zoom))
        for tag, style in self._rich_styles.items():
            self._configure_style_tag(tag, style)
        self.status_var.set(f"正文阅读缩放已调整为 {round(self._view_zoom * 100)}%（不改变保存字号）")

    def _serialize_rich_payload(self) -> Optional[dict[str, Any]]:
        styles: list[dict[str, Any]] = []
        for tag, style in self._rich_styles.items():
            raw_ranges = self.body_text.tag_ranges(tag)
            ranges: list[list[int]] = []
            for index in range(0, len(raw_ranges), 2):
                start = int(self.body_text.count("1.0", str(raw_ranges[index]), "chars")[0])
                end = int(self.body_text.count("1.0", str(raw_ranges[index + 1]), "chars")[0])
                if end > start:
                    ranges.append([start, end])
            if ranges:
                style_data = asdict(style)
                if style_data.get("baseline") == "normal":
                    style_data.pop("baseline", None)
                styles.append({"style": style_data, "ranges": ranges})
        return {"version": 1, "styles": styles} if styles else None

    @staticmethod
    def _encode_payload(payload: dict[str, Any]) -> str:
        raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        return base64.urlsafe_b64encode(raw).decode("ascii").rstrip("=")

    @staticmethod
    def _decode_rich_body(body: str) -> tuple[str, Optional[dict[str, Any]]]:
        marker_index = body.rfind(RICH_MARKER)
        if marker_index < 0 or not body.endswith(RICH_END):
            return body, None
        token = body[marker_index + len(RICH_MARKER) : -len(RICH_END)]
        try:
            token += "=" * (-len(token) % 4)
            payload = json.loads(base64.urlsafe_b64decode(token.encode("ascii")).decode("utf-8"))
            if payload.get("version") != 1:
                return body, None
            return body[:marker_index], payload
        except (binascii.Error, ValueError, TypeError, UnicodeDecodeError):
            return body, None

    @staticmethod
    def _semantic_rich_text(body: str, payload: Optional[dict[str, Any]]) -> str:
        """Render rich baseline ranges as Unicode for search and graph labels."""
        if not payload or not body:
            return body
        baselines = ["normal"] * len(body)
        for item in payload.get("styles", []):
            style = item.get("style", {}) if isinstance(item, dict) else {}
            baseline = style.get("baseline", "normal") if isinstance(style, dict) else "normal"
            if baseline not in {"super", "sub"}:
                continue
            for raw_range in item.get("ranges", []):
                try:
                    start, end = int(raw_range[0]), int(raw_range[1])
                except (IndexError, TypeError, ValueError):
                    continue
                start = min(len(body), max(0, start))
                end = min(len(body), max(start, end))
                for index in range(start, end):
                    baselines[index] = baseline
        rendered: list[str] = []
        for char, baseline in zip(body, baselines):
            if baseline == "super":
                rendered.append(char.translate(SUPERSCRIPT_MAP))
            elif baseline == "sub":
                rendered.append(char.translate(SUBSCRIPT_MAP))
            else:
                rendered.append(char)
        return "".join(rendered)

    def _semantic_editor_body(self) -> str:
        body = self.body_text.get("1.0", "end-1c")
        return self._semantic_rich_text(body, self._serialize_rich_payload())

    def _load_rich_payload(self, payload: Optional[dict[str, Any]]) -> None:
        for tag in tuple(self._rich_styles):
            self.body_text.tag_delete(tag)
        self._rich_styles.clear()
        if not payload:
            return
        body_length = len(self.body_text.get("1.0", "end-1c"))
        for item in payload.get("styles", []):
            try:
                style_data = dict(item["style"])
                style_data["size"] = min(72, max(8, int(style_data.get("size", 11))))
                if style_data.get("baseline", "normal") not in {"normal", "super", "sub"}:
                    style_data["baseline"] = "normal"
                style = RichStyle(**style_data)
                tag = self._style_tag(style)
                for start, end in item.get("ranges", []):
                    start = min(body_length, max(0, int(start)))
                    end = min(body_length, max(start, int(end)))
                    if end > start:
                        self.body_text.tag_add(tag, f"1.0+{start}c", f"1.0+{end}c")
            except (KeyError, TypeError, ValueError):
                continue

    def _compose_current_text(self) -> str:
        title = self.title_var.get().strip()
        body = self.body_text.get("1.0", "end-1c").rstrip()
        payload = self._serialize_rich_payload()
        if payload:
            body += RICH_MARKER + self._encode_payload(payload) + RICH_END
        return self._compose_text(title, body)

    @staticmethod
    def _clean_relation_entity(value: str) -> str:
        value = re.sub(r"\s+", " ", value).strip(" ，。；;、:：()（）[]【】")
        value = re.sub(r"^(?:研究|探究|考察|比较|测定|记录|分析|观察|添加|加入|采用|利用|通过|不同浓度的|不同的)", "", value)
        value = re.sub(r"(?:溶液|样品)$", lambda match: match.group(0), value)
        return value.strip(" 的")

    @staticmethod
    def _normalize_relation_label(value: str) -> str:
        label = re.sub(r"\s+", "", value).strip("，。；;、:：")
        return RELATION_LABEL_ALIASES.get(label, label)

    @classmethod
    def _combine_relation_label(cls, first: str, second: str) -> str:
        first = cls._normalize_relation_label(first)
        second = cls._normalize_relation_label(second)
        combined = f"{first}{second}"
        return RELATION_LABEL_ALIASES.get(combined, combined)

    @classmethod
    def _extract_relations_from_entries(
        cls,
        entries: list[dict[str, Any]],
        keyword_rules: Optional[list[KeywordRule]] = None,
    ) -> list[Relation]:
        alias_map = dict(RELATION_LABEL_ALIASES)
        relation_importance: dict[str, int] = {}
        concept_aliases: dict[str, str] = {}
        concept_importance: dict[str, int] = {}
        for rule in keyword_rules or []:
            if not rule.enabled:
                continue
            if rule.kind == "relation":
                canonical = rule.display_label
                alias_map[rule.term] = canonical
                alias_map.setdefault(canonical, canonical)
                relation_importance[canonical.casefold()] = max(
                    rule.importance, relation_importance.get(canonical.casefold(), 1)
                )
            else:
                concept_aliases[rule.term] = rule.display_label
                concept_importance[rule.display_label.casefold()] = max(
                    rule.importance, concept_importance.get(rule.display_label.casefold(), 1)
                )
        relation_words = tuple(sorted(alias_map, key=len, reverse=True))
        relation_pattern = "|".join(map(re.escape, relation_words))
        entity_char = r"[A-Za-z0-9α-ωΑ-Ω\u4e00-\u9fff\u2070-\u209f¹²³·（）()\-⁺⁻₊₋]"
        entity = rf"{entity_char}(?:[ \t]*{entity_char}){{0,31}}?"
        boundary = r"(?:^|[，。；;、\n：:])\s*"
        ending = r"(?=\s*(?:实验|过程|作用)?\s*(?:$|[，。；;、\n：:]))"
        pair_patterns = (
            re.compile(
                rf"{boundary}(?P<a>{entity})\s*(?:与|和|及|\+)\s*(?P<b>{entity})\s*(?:之间)?\s*(?:发生|进行|产生|的)?\s*(?P<rel>{relation_pattern})(?:实验|过程)?",
                re.IGNORECASE,
            ),
            re.compile(
                rf"{boundary}(?P<a>{entity})\s*对\s*(?P<b>{entity})\s*(?:的|发生|进行|产生)?\s*(?P<rel>{relation_pattern})(?:实验|过程)?",
                re.IGNORECASE,
            ),
            re.compile(
                rf"{boundary}(?P<b>{entity})\s*(?:被|受)\s*(?P<a>{entity})\s*(?P<rel>{relation_pattern})(?:实验|过程)?",
                re.IGNORECASE,
            ),
        )
        direct_pattern = re.compile(
            rf"{boundary}(?P<a>{entity})\s*(?P<rel>{relation_pattern})\s*(?:了|着|对|使|使得|为)?\s*(?P<b>{entity}){ending}",
            re.IGNORECASE,
        )
        coupled_pattern = re.compile(
            rf"{boundary}(?P<a>{entity})\s*(?P<first>促进|抑制|增强|削弱|加快|减慢)\s*(?P<b>{entity})\s*(?:的)?\s*(?P<second>{relation_pattern}){ending}",
            re.IGNORECASE,
        )
        catalysis_pattern = re.compile(
            rf"{boundary}(?P<a>{entity})\s*(?P<first>光催化|电催化|生物催化|催化)\s*(?P<b>{entity})\s*(?:发生|进行)?\s*(?P<second>降解|分解|转化|氧化|还原|矿化){ending}",
            re.IGNORECASE,
        )
        under_action_pattern = re.compile(
            rf"{boundary}(?P<b>{entity})\s*在\s*(?P<a>{entity})\s*(?:的)?\s*(?P<mode>催化|作用)下\s*(?:发生|进行|被)?\s*(?P<rel>{relation_pattern}){ending}",
            re.IGNORECASE,
        )
        experiment_base_pattern = re.compile(
            rf"{boundary}(?P<a>实验{entity})\s*(?:基于|以)\s*(?P<b>实验{entity})\s*(?:为基础|的基础上)?\s*(?:进行|开展|完成)?\s*(?P<rel>优化|改进|验证|延续|衍生|复现){ending}",
            re.IGNORECASE,
        )
        experiment_role_pattern = re.compile(
            rf"{boundary}(?P<a>实验{entity})\s*(?:是|作为)\s*(?P<b>实验{entity})\s*(?:的)?\s*(?P<rel>对照|前置|后续)(?:实验)?{ending}",
            re.IGNORECASE,
        )
        trend_pattern = re.compile(
            rf"(?:^|[。；;、\n：:])\s*随着\s*(?P<a>[^，,。；;\n]{{1,18}})\s*[，,]\s*(?P<b>{entity})\s*(?P<rel>提高|提升|增大|增加|降低|减小|减少|加快|减慢|增强|削弱){ending}",
            re.IGNORECASE,
        )
        relations: list[Relation] = []
        seen: set[tuple[str, str, str, int]] = set()

        def normalize_label(value: str) -> str:
            cleaned = re.sub(r"\s+", "", value).strip("，。；;、:：")
            return alias_map.get(cleaned, cleaned)

        def combine_label(first: str, second: str) -> str:
            first = normalize_label(first)
            second = normalize_label(second)
            combined = f"{first}{second}"
            return alias_map.get(combined, combined)

        def add_relation(source: str, target: str, label: str, entry_id: int) -> None:
            source = cls._clean_relation_entity(source)
            target = cls._clean_relation_entity(target)
            label = normalize_label(label)
            if not source or not target or source == target or len(source) > 28 or len(target) > 28:
                return
            key = (source.casefold(), target.casefold(), label, entry_id)
            if key in seen:
                return
            seen.add(key)
            importance = max(
                relation_importance.get(label.casefold(), 1),
                concept_importance.get(source.casefold(), 1),
                concept_importance.get(target.casefold(), 1),
            )
            relations.append(Relation(source, target, label, entry_id, importance))

        for entry in entries:
            try:
                entry_id = int(entry.get("id"))
            except (TypeError, ValueError):
                continue
            title, body, rich_payload = cls._decode_entry_text(entry.get("text", ""))
            body = cls._semantic_rich_text(body, rich_payload)
            text = f"{title}。{body}".replace("＋", "+")
            for term in sorted(concept_aliases, key=len, reverse=True):
                text = re.sub(re.escape(term), lambda _match, replacement=concept_aliases[term]: replacement, text, flags=re.IGNORECASE)
            for pattern in pair_patterns:
                for match in pattern.finditer(text):
                    add_relation(match.group("a"), match.group("b"), match.group("rel"), entry_id)
            for match in catalysis_pattern.finditer(text):
                source = cls._clean_relation_entity(match.group("a"))
                target = cls._clean_relation_entity(match.group("b"))
                if ("在" in source and target == "下") or target in {"下", "作用下", "催化下"}:
                    continue
                label = combine_label(match.group("first"), match.group("second"))
                add_relation(source, target, label, entry_id)
            for match in coupled_pattern.finditer(text):
                label = combine_label(match.group("first"), match.group("second"))
                add_relation(match.group("a"), match.group("b"), label, entry_id)
            for match in under_action_pattern.finditer(text):
                label = match.group("rel")
                if match.group("mode") == "催化":
                    label = combine_label("催化", label)
                add_relation(match.group("a"), match.group("b"), label, entry_id)
            for pattern in (experiment_base_pattern, experiment_role_pattern, trend_pattern):
                for match in pattern.finditer(text):
                    add_relation(match.group("a"), match.group("b"), match.group("rel"), entry_id)
            for match in direct_pattern.finditer(text):
                source = cls._clean_relation_entity(match.group("a"))
                target = cls._clean_relation_entity(match.group("b"))
                if re.search(r"(?:与|和|及|被|受|之间|发生|进行|产生)", source) or source.endswith("对"):
                    continue
                if target in {"实验", "过程", "作用", "反应"} or any(target.endswith(word) for word in relation_words):
                    continue
                add_relation(source, target, match.group("rel"), entry_id)
        return relations

    @staticmethod
    def _relation_node_kind(value: str) -> str:
        compact = re.sub(r"\s+", "", value)
        if compact.startswith("实验") or compact.endswith("实验"):
            return "experiment"
        condition_hints = (
            "pH", "温度", "压力", "浓度", "时间", "流量", "转速", "剂量", "投加量",
            "去除率", "转化率", "选择性", "产率", "速率", "吸光度", "效率", "压降",
        )
        if any(hint.casefold() in compact.casefold() for hint in condition_hints):
            return "condition"
        return "substance"

    @staticmethod
    def _display_node_label(value: str) -> str:
        compact = re.sub(r"\s+", " ", value).strip()
        return compact if len(compact) <= 12 else compact[:11] + "…"

    def _entries_for_relation_analysis(self) -> list[dict[str, Any]]:
        entries = [dict(entry) for entry in (self.all_entries or self.entries)]
        if not self.current:
            return entries
        try:
            current_id = int(self.current.get("id"))
        except (TypeError, ValueError):
            return entries
        live_text = self._compose_text(
            self.title_var.get(),
            self._semantic_editor_body(),
        )
        attachment_parts: list[str] = []
        for attachment in self.current.get("attachments", []) or []:
            try:
                attachment_id = int(attachment.get("id"))
            except (TypeError, ValueError):
                continue
            extracted = self._attachment_text_cache.get(attachment_id, "")
            if extracted:
                attachment_parts.append(extracted)
        if attachment_parts:
            live_text += "\n" + "\n".join(attachment_parts)
        live_entry = {**self.current, "id": current_id, "text": live_text}
        for index, entry in enumerate(entries):
            try:
                entry_id = int(entry.get("id"))
            except (TypeError, ValueError):
                continue
            if entry_id == current_id:
                entries[index] = live_entry
                break
        else:
            entries.append(live_entry)
        return entries

    def _rebuild_relation_graph(self) -> None:
        if self._relation_analysis_job:
            self.after_cancel(self._relation_analysis_job)
            self._relation_analysis_job = None
        self._relations = self._extract_relations_from_entries(self._entries_for_relation_analysis(), self._keyword_rules)
        self._draw_relation_graph()

    def _schedule_relation_analysis(self) -> None:
        if self._relation_analysis_job:
            self.after_cancel(self._relation_analysis_job)
        self._relation_analysis_job = self.after(420, self._rebuild_relation_graph)

    def _schedule_graph_draw(self, _event: Optional[tk.Event] = None) -> None:
        if self._graph_job:
            self.after_cancel(self._graph_job)
        self._graph_job = self.after(80, self._draw_relation_graph)

    @staticmethod
    def _zoom_pan_for_anchor(
        old_scale: float,
        new_scale: float,
        pan_x: float,
        pan_y: float,
        anchor_x: float,
        anchor_y: float,
        center_x: float,
        center_y: float,
    ) -> tuple[float, float]:
        ratio = new_scale / max(old_scale, 0.0001)
        new_pan_x = anchor_x - center_x - (anchor_x - center_x - pan_x) * ratio
        new_pan_y = anchor_y - center_y - (anchor_y - center_y - pan_y) * ratio
        return new_pan_x, new_pan_y

    def _on_graph_mousewheel(self, event: tk.Event) -> str:
        direction = 0
        if getattr(event, "delta", 0):
            direction = 1 if event.delta > 0 else -1
        elif getattr(event, "num", 0) in (4, 5):
            direction = 1 if event.num == 4 else -1
        if not direction or not self.current:
            return "break"
        old_scale = self._graph_spacing
        factor = GRAPH_SPACING_STEP if direction > 0 else 1.0 / GRAPH_SPACING_STEP
        new_scale = min(MAX_GRAPH_SPACING, max(MIN_GRAPH_SPACING, old_scale * factor))
        if abs(new_scale - old_scale) < 0.0001:
            return "break"
        width = max(10, self.mindmap_canvas.winfo_width())
        height = max(10, self.mindmap_canvas.winfo_height())
        self._graph_pan_x, self._graph_pan_y = self._zoom_pan_for_anchor(
            old_scale,
            new_scale,
            self._graph_pan_x,
            self._graph_pan_y,
            float(event.x),
            float(event.y),
            width / 2,
            height / 2,
        )
        self._graph_spacing = new_scale
        self._draw_relation_graph()
        return "break"

    def _on_graph_drag_start(self, event: tk.Event) -> str:
        if not self.current:
            return "break"
        self._graph_drag_anchor = (event.x, event.y)
        self.mindmap_canvas.configure(cursor="fleur")
        return "break"

    def _on_graph_drag_move(self, event: tk.Event) -> str:
        if self._graph_drag_anchor is None:
            return "break"
        old_x, old_y = self._graph_drag_anchor
        dx, dy = event.x - old_x, event.y - old_y
        if dx or dy:
            self._graph_pan_x += dx
            self._graph_pan_y += dy
            self.mindmap_canvas.move("all", dx, dy)
            self._graph_drag_anchor = (event.x, event.y)
        return "break"

    def _on_graph_drag_end(self, _event: tk.Event) -> str:
        self._graph_drag_anchor = None
        self.mindmap_canvas.configure(cursor="")
        return "break"

    def _reset_graph_view(self, redraw: bool = True) -> str:
        self._graph_spacing = 1.0
        self._graph_pan_x = 0.0
        self._graph_pan_y = 0.0
        self._graph_drag_anchor = None
        if redraw and hasattr(self, "mindmap_canvas"):
            self._draw_relation_graph()
        return "break"

    def _apply_graph_view(
        self,
        positions: dict[str, tuple[float, float]],
        width: float,
        height: float,
    ) -> dict[str, tuple[float, float]]:
        cx, cy = width / 2, height / 2
        return {
            node: (
                cx + (x - cx) * self._graph_spacing + self._graph_pan_x,
                cy + (y - cy) * self._graph_spacing + self._graph_pan_y,
            )
            for node, (x, y) in positions.items()
        }

    def _draw_relation_graph(self) -> None:
        self._graph_job = None
        canvas = self.mindmap_canvas
        canvas.delete("all")
        width = max(10, canvas.winfo_width())
        height = max(10, canvas.winfo_height())
        if not self.current:
            self._draw_graph_empty("选择一条实验记录", "系统会分析物质、条件和实验之间的关系")
            self.graph_status_var.set("选择记录后自动生成")
            return

        current_id = int(self.current.get("id", -1))
        selected_relations = [relation for relation in self._relations if relation.entry_id == current_id]
        if not selected_relations:
            self._draw_graph_empty("未识别到明确关系", "可描述物质作用、条件影响或实验间的对照与优化关系")
            self.graph_status_var.set("当前记录暂无可绘制的物质关系")
            return

        selected_nodes = [node for relation in selected_relations for node in (relation.source, relation.target)]
        degree: dict[str, int] = {}
        node_importance: dict[str, int] = {}
        for relation in selected_relations:
            node_importance[relation.source] = max(node_importance.get(relation.source, 1), relation.importance)
            node_importance[relation.target] = max(node_importance.get(relation.target, 1), relation.importance)
        for node in selected_nodes:
            degree[node] = degree.get(node, 0) + 1
        center = max(degree, key=lambda node: (node_importance.get(node, 1), degree[node], -selected_nodes.index(node)))

        nodes: list[str] = [center]
        for relation in selected_relations:
            for node in (relation.source, relation.target):
                if node not in nodes:
                    nodes.append(node)
        for relation in self._relations:
            if len(nodes) >= 9:
                break
            if relation.source in nodes and relation.target not in nodes:
                nodes.append(relation.target)
            elif relation.target in nodes and relation.source not in nodes:
                nodes.append(relation.source)

        visible = set(nodes)
        edges: list[Relation] = []
        edge_keys: set[tuple[str, str, str]] = set()
        for relation in selected_relations + self._relations:
            if relation.source not in visible or relation.target not in visible:
                continue
            key = tuple(sorted((relation.source, relation.target))) + (relation.label,)
            if key in edge_keys:
                continue
            edge_keys.add(key)
            edges.append(relation)

        node_radius = max(34, min(48, min(width, height) * 0.095))

        adjacency: dict[str, set[str]] = {node: set() for node in nodes}
        for relation in edges:
            adjacency[relation.source].add(relation.target)
            adjacency[relation.target].add(relation.source)
        components: list[list[str]] = []
        unseen = set(nodes)
        while unseen:
            seed = next(node for node in nodes if node in unseen)
            stack = [seed]
            component: list[str] = []
            unseen.remove(seed)
            while stack:
                node = stack.pop()
                component.append(node)
                for neighbor in adjacency[node]:
                    if neighbor in unseen:
                        unseen.remove(neighbor)
                        stack.append(neighbor)
            components.append(component)

        positions: dict[str, tuple[float, float]] = {}
        primary_node = center if len(components) == 1 else ""
        if len(components) == 1:
            cx, cy = width / 2, height / 2
            radius = max(82, min(width, height) * (0.29 if len(nodes) <= 6 else 0.34))
            positions[center] = (cx, cy)
            others = [node for node in nodes if node != center]
            for index, node in enumerate(others):
                angle = -math.pi / 2 + (2 * math.pi * index / max(1, len(others)))
                positions[node] = (cx + radius * math.cos(angle), cy + radius * math.sin(angle))
        else:
            columns = 2 if width >= 620 and len(components) >= 4 else 1
            rows = math.ceil(len(components) / columns)
            cell_width = width / columns
            cell_height = height / rows
            for component_index, component in enumerate(components):
                column = component_index % columns
                row = component_index // columns
                cx = cell_width * (column + 0.5)
                cy = cell_height * (row + 0.5)
                if len(component) == 1:
                    positions[component[0]] = (cx, cy)
                    continue
                if len(component) == 2:
                    gap = min(max(node_radius * 2.35, 96), cell_width * 0.32)
                    positions[component[0]] = (cx - gap, cy)
                    positions[component[1]] = (cx + gap, cy)
                    continue
                component_center = max(component, key=lambda node: len(adjacency[node]))
                positions[component_center] = (cx, cy)
                others = [node for node in component if node != component_center]
                component_radius = min(cell_width, cell_height) * 0.28
                component_radius = max(node_radius * 2.2, component_radius)
                for index, node in enumerate(others):
                    angle = -math.pi / 2 + (2 * math.pi * index / len(others))
                    positions[node] = (
                        cx + component_radius * math.cos(angle),
                        cy + component_radius * math.sin(angle),
                    )

        positions = self._apply_graph_view(positions, width, height)

        for relation in edges:
            node_importance[relation.source] = max(node_importance.get(relation.source, 1), relation.importance)
            node_importance[relation.target] = max(node_importance.get(relation.target, 1), relation.importance)
            x1, y1 = positions[relation.source]
            x2, y2 = positions[relation.target]
            dx, dy = x2 - x1, y2 - y1
            distance = max(1.0, math.hypot(dx, dy))
            ux, uy = dx / distance, dy / distance
            sx, sy = x1 + ux * node_radius, y1 + uy * node_radius
            ex, ey = x2 - ux * node_radius, y2 - uy * node_radius
            line_color = C.primary if relation.importance >= 3 else C.border_strong
            canvas.create_line(sx, sy, ex, ey, fill=line_color, width=self._px(1 + relation.importance), smooth=True)
            mx, my = (sx + ex) / 2, (sy + ey) / 2
            label = canvas.create_text(
                mx,
                my - self._px(6),
                text=relation.label,
                fill=C.text_muted,
                font=self._font(9, "bold" if relation.importance >= 2 else "normal"),
                anchor="s",
            )
            box = canvas.bbox(label)
            if box:
                background = canvas.create_rectangle(
                    box[0] - 4,
                    box[1] - 2,
                    box[2] + 4,
                    box[3] + 2,
                    fill=C.surface,
                    outline="",
                )
                canvas.tag_lower(background, label)

        current_entities = set(selected_nodes)
        for node in reversed(nodes):
            x, y = positions[node]
            is_center = node == primary_node
            is_current = node in current_entities
            importance = node_importance.get(node, 1)
            node_kind = self._relation_node_kind(node)
            if is_center:
                fill, outline, text_color = C.primary, C.primary, "white"
            elif node_kind == "experiment":
                fill, outline, text_color = C.success_soft, C.success, C.navy
            elif node_kind == "condition":
                fill, outline, text_color = C.warning_soft, C.warning, C.navy
            else:
                fill = C.primary_soft if is_current else C.surface_alt
                outline = C.primary if is_current else C.border_strong
                text_color = C.navy
            canvas.create_oval(
                x - node_radius,
                y - node_radius,
                x + node_radius,
                y + node_radius,
                fill=fill,
                outline=outline,
                width=self._px(max(2 if is_current else 1, importance)),
            )
            canvas.create_text(
                x,
                y,
                text=self._display_node_label(node),
                fill=text_color,
                font=self._font(10, "bold" if is_current else "normal"),
                width=max(40, round(node_radius * 1.65)),
                justify="center",
            )

        if len(components) == 1:
            self.graph_status_var.set(
                f"实时预览 · 以“{center}”为中心 · {len(nodes)} 个节点 · {len(edges)} 条关系 · 间距{round(self._graph_spacing * 100)}%"
            )
        else:
            self.graph_status_var.set(
                f"实时预览 · {len(components)} 组关系 · {len(nodes)} 个节点 · {len(edges)} 条连线 · 间距{round(self._graph_spacing * 100)}%"
            )

    def _draw_graph_empty(self, title: str, subtitle: str) -> None:
        canvas = self.mindmap_canvas
        width = max(10, canvas.winfo_width())
        height = max(10, canvas.winfo_height())
        cx, cy = width / 2, height / 2
        radius = self._px(34)
        canvas.create_oval(cx - radius, cy - radius - 24, cx + radius, cy + radius - 24, fill=C.primary_soft, outline=C.primary, width=self._px(2))
        canvas.create_text(cx, cy - 24, text="关系", fill=C.primary, font=self._font(10, "bold"))
        canvas.create_text(cx, cy + 34, text=title, fill=C.text, font=self._font(12, "bold"))
        canvas.create_text(cx, cy + 60, text=subtitle, fill=C.text_muted, font=self._font(9.5), width=max(160, width - 50), justify="center")

    def _focus_search(self) -> str:
        self.search_entry.focus_set()
        self.search_entry.select_range(0, tk.END)
        return "break"

    def _set_connected(self, connected: bool, label: str) -> None:
        self.connection_var.set(label)
        self.connection_dot.configure(fg="#5CD09A" if connected else "#F1B451")

    def _set_busy(self, active: bool, message: str = "") -> None:
        self._busy_count = max(0, self._busy_count + (1 if active else -1))
        if message:
            self.status_var.set(message)
        if self._busy_count:
            self.configure(cursor="watch")
        else:
            self.configure(cursor="")

    def _run(self, description: str, work: Callable[[], Any], on_ok: Optional[Callable[[Any], None]] = None) -> None:
        self._set_busy(True, f"{description}…")

        def worker() -> None:
            try:
                result = work()
                self._queue.put(("ok", (description, on_ok), result))
            except Exception as exc:  # network and API errors are surfaced in the UI
                self._queue.put(("error", description, exc))

        threading.Thread(target=worker, daemon=True).start()

    def _poll_queue(self) -> None:
        try:
            while True:
                kind, context, payload = self._queue.get_nowait()
                self._set_busy(False)
                if kind == "ok":
                    description, callback = context
                    self.status_var.set(f"{description}完成")
                    if callback:
                        callback(payload)
                else:
                    self.status_var.set(f"{context}失败")
                    messagebox.showerror(f"{context}失败", str(payload), parent=self)
        except queue.Empty:
            pass
        if not self._closing:
            self.after(80, self._poll_queue)

    def _connect(self) -> None:
        if self._busy_count:
            return
        self._set_connected(False, "正在连接")

        def create_client() -> Any:
            if self.demo:
                client = DemoClient()
            else:
                from config import API_TOKEN, REMOTE_HOST, REMOTE_PORT, SSH_HOST, SSH_PORT, SSH_PRIVATE_KEY, SSH_USER
                from ssh_http import ChemDBClient

                client = ChemDBClient(
                    ssh_host=SSH_HOST,
                    ssh_port=SSH_PORT,
                    ssh_user=SSH_USER,
                    private_key_str=SSH_PRIVATE_KEY,
                    remote_host=REMOTE_HOST,
                    remote_port=REMOTE_PORT,
                    api_token=API_TOKEN,
                )
            client.connect()
            client.health()
            return client

        def connected(client: Any) -> None:
            old_client = self.client
            self.client = client
            if old_client and old_client is not client:
                try:
                    old_client.close()
                except Exception:
                    pass
            self._set_connected(True, "演示模式" if self.demo else "服务器已连接")
            self.refresh()

        self._run("连接服务器", create_client, connected)

    def _schedule_search(self, _event: tk.Event) -> None:
        if self._search_job:
            self.after_cancel(self._search_job)
        self._search_job = self.after(380, self.refresh)

    def _show_all(self) -> None:
        self.search_var.set("")
        self.refresh()

    def refresh(self, select_id: Optional[int] = None) -> None:
        if not self.client:
            return
        query = self.search_var.get().strip()
        if select_id is None and self.current:
            select_id = self.current.get("id")

        def load() -> Any:
            visible = self.client.list_entries(q=query, limit=200, offset=0)
            all_rows = visible if not query else self.client.list_entries(q="", limit=200, offset=0)
            return {"visible": visible, "all": all_rows}

        self._run("查询记录", load, lambda rows: self._fill_list(rows, select_id))

    def _fill_list(self, entries: Any, select_id: Optional[int] = None) -> None:
        all_entries: Any = entries
        if isinstance(entries, dict) and "visible" in entries:
            all_entries = entries.get("all", [])
            entries = entries.get("visible", [])
        if isinstance(entries, dict):
            entries = entries.get("items", entries.get("entries", []))
        if isinstance(all_entries, dict):
            all_entries = all_entries.get("items", all_entries.get("entries", []))
        self.entries = list(entries or [])
        self.all_entries = list(all_entries or [])
        self._relations = self._extract_relations_from_entries(self.all_entries, self._keyword_rules)
        self.records.delete(*self.records.get_children())
        selected_iid: Optional[str] = None
        for entry in self.entries:
            title, _body = self._split_text(entry.get("text", ""))
            title = title or "未命名实验记录"
            count = len(entry.get("attachments", []))
            iid = str(entry["id"])
            self.records.insert("", "end", iid=iid, values=(title, str(count)))
            if entry.get("id") == select_id:
                selected_iid = iid
        self.count_var.set(f"{len(self.entries)} 条记录")
        if selected_iid:
            self.records.selection_set(selected_iid)
            self.records.focus(selected_iid)
            self.records.see(selected_iid)
            self.after_idle(lambda: self.records.event_generate("<<TreeviewSelect>>"))
        elif self.entries:
            first_iid = str(self.entries[0]["id"])
            self.records.selection_set(first_iid)
            self.records.focus(first_iid)
            self.records.see(first_iid)
            self.after_idle(lambda: self.records.event_generate("<<TreeviewSelect>>"))
        elif not self.entries:
            self.current = None
            self._show_empty_state("没有找到匹配记录", "换个关键词试试，或新建一条实验记录")
        self._draw_relation_graph()

    def _on_select(self, _event: tk.Event) -> None:
        if self._restoring_selection:
            return
        selection = self.records.selection()
        if not selection or not self.client:
            return
        entry_id = int(selection[0])
        if self.current and entry_id == int(self.current.get("id", -1)):
            return
        if self._dirty and self.current:
            discard = messagebox.askyesno(
                "有未保存修改",
                "当前记录的文字或格式尚未保存。\n\n是否放弃修改并切换记录？",
                icon="warning",
                parent=self,
            )
            if not discard:
                current_iid = str(self.current.get("id"))
                if self.records.exists(current_iid):
                    self._restoring_selection = True
                    self.records.selection_set(current_iid)
                    self.records.focus(current_iid)
                    self.after_idle(self._finish_restoring_selection)
                return
        self._run("加载记录", lambda: self.client.get_entry(entry_id), self._show_detail)

    def _finish_restoring_selection(self) -> None:
        self._restoring_selection = False

    def _show_detail(self, entry: dict[str, Any]) -> None:
        self._updating_editor = True
        previous_id = self.current.get("id") if self.current else None
        self.current = entry
        if previous_id != entry.get("id"):
            self._reset_graph_view(redraw=False)
        title, body, payload = self._decode_entry_text(entry.get("text", ""))
        self.title_var.set(title)
        self.body_text.configure(state="normal")
        self.body_text.delete("1.0", tk.END)
        self.body_text.insert("1.0", body)
        self._load_rich_payload(payload)
        self.body_text.edit_modified(False)
        self._saved_text = self._compose_current_text()
        self._dirty = False
        self._updating_editor = False

        attachments = entry.get("attachments", []) or []
        self.attachments.delete(*self.attachments.get_children())
        for item in attachments:
            file_type = self._file_type(item.get("filename", ""))
            size = self._format_size(item.get("size", 0))
            self.attachments.insert("", "end", iid=str(item["id"]), values=(item.get("filename", "未命名文件"), file_type, size))

        self._update_selection_meta()
        self.save_button.configure(text="保存修改  Ctrl+S")
        self.preview_hint_var.set("选择附件后可在此预览")
        self.preview_label.configure(image="")
        self._photo = None
        self.empty_state.place_forget()
        self._draw_relation_graph()
        self._load_table_attachments(entry)

        for item in attachments:
            if item.get("has_thumb"):
                self.attachments.selection_set(str(item["id"]))
                self._fetch_thumb(item["id"])
                break

    @staticmethod
    def _extract_table_text(filename: str, data: bytes, max_cells: int = 6000, max_chars: int = 160000) -> str:
        ext = Path(filename).suffix.lower()
        rows: list[list[str]] = []
        cell_count = 0

        def add_row(values: Any) -> bool:
            nonlocal cell_count
            cells = ["" if value is None else str(value).strip() for value in values]
            if any(cells):
                remaining = max_cells - cell_count
                if remaining <= 0:
                    return False
                rows.append(cells[:remaining])
                cell_count += min(len(cells), remaining)
            return cell_count < max_cells

        if ext in {".csv", ".tsv"}:
            decoded = ""
            for encoding in ("utf-8-sig", "gb18030", "utf-16"):
                try:
                    decoded = data.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if not decoded:
                decoded = data.decode("utf-8", errors="replace")
            delimiter = "\t" if ext == ".tsv" else ","
            try:
                dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t|")
                delimiter = dialect.delimiter
            except csv.Error:
                pass
            for row in csv.reader(io.StringIO(decoded), delimiter=delimiter):
                if not add_row(row):
                    break
        elif ext == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("读取 XLSX 需要安装 openpyxl") from exc
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    add_row([f"工作表 {sheet.title}"])
                    for row in sheet.iter_rows(values_only=True):
                        if not add_row(row):
                            break
                    if cell_count >= max_cells:
                        break
            finally:
                workbook.close()
        elif ext == ".xls":
            try:
                import xlrd
            except ImportError as exc:
                raise RuntimeError("读取旧版 XLS 需要安装 xlrd") from exc
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                for sheet in workbook.sheets():
                    add_row([f"工作表 {sheet.name}"])
                    for row_index in range(sheet.nrows):
                        if not add_row(sheet.row_values(row_index)):
                            break
                    if cell_count >= max_cells:
                        break
            finally:
                workbook.release_resources()
        else:
            return ""
        relation_sentences: list[str] = []
        for header_index, header in enumerate(rows):
            normalized = [re.sub(r"\s+", "", cell).casefold() for cell in header]

            def column(names: set[str]) -> Optional[int]:
                return next((index for index, cell in enumerate(normalized) if cell in names), None)

            source_col = column({"物质a", "反应物a", "反应物1", "主体", "来源", "source"})
            target_col = column({"物质b", "反应物b", "反应物2", "客体", "对象", "target"})
            relation_col = column({"关系", "关系词", "反应", "反应类型", "作用", "relation"})
            if source_col is None or target_col is None or relation_col is None:
                continue
            max_column = max(source_col, target_col, relation_col)
            for row in rows[header_index + 1 :]:
                if len(row) <= max_column:
                    continue
                source, target, relation = row[source_col], row[target_col], row[relation_col]
                if source and target and relation:
                    relation_sentences.append(f"{source}与{target}发生{relation}。")
            break
        plain_text = "\n".join(" ".join(row) for row in rows)
        if relation_sentences:
            plain_text += "\n" + "\n".join(relation_sentences)
        return plain_text[:max_chars]

    def _load_table_attachments(self, entry: dict[str, Any]) -> None:
        if not self.client:
            return
        pending: list[tuple[int, str]] = []
        for attachment in entry.get("attachments", []) or []:
            filename = str(attachment.get("filename", ""))
            if Path(filename).suffix.lower() not in {".csv", ".tsv", ".xlsx", ".xls"}:
                continue
            try:
                attachment_id = int(attachment.get("id"))
            except (TypeError, ValueError):
                continue
            if attachment_id not in self._attachment_text_cache and attachment_id not in self._attachment_parse_errors:
                pending.append((attachment_id, filename))
        if not pending:
            return
        try:
            entry_id = int(entry.get("id"))
        except (TypeError, ValueError):
            return

        def read_tables() -> tuple[dict[int, str], dict[int, str]]:
            extracted: dict[int, str] = {}
            errors: dict[int, str] = {}
            for attachment_id, filename in pending:
                try:
                    data = self.client.download_attachment(attachment_id)
                    if len(data) > 25 * 1024 * 1024:
                        raise ValueError("表格超过 25 MB，已跳过自动分析")
                    extracted[attachment_id] = self._extract_table_text(filename, data)
                except Exception as exc:
                    errors[attachment_id] = str(exc)
            return extracted, errors

        def loaded(result: tuple[dict[int, str], dict[int, str]]) -> None:
            extracted, errors = result
            self._attachment_text_cache.update(extracted)
            self._attachment_parse_errors.update(errors)
            if self.current and int(self.current.get("id", -1)) == entry_id:
                self._rebuild_relation_graph()
                ok_count = sum(1 for value in extracted.values() if value)
                if ok_count:
                    self.status_var.set(f"已提取 {ok_count} 个表格附件并更新思维导图")
                elif errors:
                    self.status_var.set("表格附件未能解析；请检查格式或读取组件")

        self._run("读取表格附件", read_tables, loaded)

    def _show_empty_state(self, title: str = "选择一条实验记录", subtitle: str = "从左侧打开记录，或新建一条实验记录") -> None:
        self._dirty = False
        self._saved_text = ""
        labels = [w for w in self.empty_state.winfo_children()[0].winfo_children() if isinstance(w, ttk.Label)]
        if len(labels) >= 2:
            labels[0].configure(text=title)
            labels[1].configure(text=subtitle)
        self.empty_state.place(relx=0, rely=0, relwidth=1, relheight=1)
        self.empty_state.lift()
        if hasattr(self, "mindmap_canvas"):
            self._draw_relation_graph()

    @staticmethod
    def _split_text(text: str) -> tuple[str, str]:
        lines = (text or "").splitlines()
        if not lines:
            return "", ""
        body = "\n".join(lines[1:]).lstrip("\n")
        body, _payload = ChemDBApp._decode_rich_body(body)
        return lines[0].strip(), body

    @staticmethod
    def _decode_entry_text(text: str) -> tuple[str, str, Optional[dict[str, Any]]]:
        lines = (text or "").splitlines()
        if not lines:
            return "", "", None
        body = "\n".join(lines[1:]).lstrip("\n")
        body, payload = ChemDBApp._decode_rich_body(body)
        return lines[0].strip(), body, payload

    @staticmethod
    def _compose_text(title: str, body: str) -> str:
        title = title.strip()
        body = body.strip()
        return f"{title}\n{body}".strip() if body else title

    def _new_entry_dialog(self) -> None:
        if not self.client:
            messagebox.showwarning("尚未连接", "请等待服务器连接成功后再新建记录。", parent=self)
            return

        dialog = tk.Toplevel(self)
        dialog.title("新建实验记录")
        dialog.geometry(f"{self._px(700)}x{self._px(640)}")
        dialog.minsize(self._px(600), self._px(540))
        dialog.configure(bg=C.canvas)
        dialog.transient(self)
        dialog.grab_set()
        dialog.grid_columnconfigure(0, weight=1)
        dialog.grid_rowconfigure(1, weight=1)

        head = ttk.Frame(dialog, style="Header.TFrame", padding=(20, 15))
        head.grid(row=0, column=0, sticky="ew")
        ttk.Label(head, text="新建实验记录", style="HeaderTitle.TLabel").pack(anchor="w")
        ttk.Label(head, text="填写标题、实验内容，并可同时上传原始数据附件", style="HeaderSub.TLabel").pack(anchor="w", pady=(3, 0))

        form = tk.Frame(dialog, bg=C.surface, highlightbackground=C.border, highlightthickness=1)
        form.grid(row=1, column=0, sticky="nsew", padx=18, pady=18)
        form.grid_columnconfigure(0, weight=1)
        form.grid_rowconfigure(3, weight=1)

        ttk.Label(form, text="实验标题 *", style="Section.TLabel").grid(row=0, column=0, sticky="w", padx=18, pady=(16, 7))
        title = ttk.Entry(form)
        title.grid(row=1, column=0, sticky="ew", padx=18)

        ttk.Label(form, text="实验内容", style="Section.TLabel").grid(row=2, column=0, sticky="w", padx=18, pady=(14, 7))
        content_host = tk.Frame(form, bg=C.surface, highlightbackground=C.border_strong, highlightthickness=1)
        content_host.grid(row=3, column=0, sticky="nsew", padx=18)
        content_host.grid_rowconfigure(0, weight=1)
        content_host.grid_columnconfigure(0, weight=1)
        content = tk.Text(content_host, wrap="word", relief="flat", padx=self._px(10), pady=self._px(9), undo=True, bg=C.surface, fg=C.text, insertbackground=C.primary, font=self._font(11))
        content.grid(row=0, column=0, sticky="nsew")

        files: list[str] = []
        file_label = tk.StringVar(value="尚未选择附件")
        files_row = ttk.Frame(form, style="Surface.TFrame", padding=(18, 12, 18, 0))
        files_row.grid(row=4, column=0, sticky="ew")
        files_row.grid_columnconfigure(0, weight=1)
        ttk.Label(files_row, textvariable=file_label, style="Muted.TLabel").grid(row=0, column=0, sticky="w")

        def choose_files() -> None:
            selected = filedialog.askopenfilenames(parent=dialog, title="选择实验附件")
            for path in selected:
                if path not in files:
                    files.append(path)
            file_label.set(f"已选择 {len(files)} 个附件" if files else "尚未选择附件")

        ttk.Button(files_row, text="选择附件", style="Secondary.TButton", command=choose_files).grid(row=0, column=1, sticky="e")

        bar = ttk.Frame(form, style="Surface.TFrame", padding=(18, 14, 18, 16))
        bar.grid(row=5, column=0, sticky="ew")
        bar.grid_columnconfigure(0, weight=1)
        ttk.Button(bar, text="取消", style="Secondary.TButton", command=dialog.destroy).grid(row=0, column=1, padx=(0, 8))

        def submit() -> None:
            title_text = title.get().strip()
            if not title_text:
                messagebox.showwarning("缺少标题", "请填写实验标题。", parent=dialog)
                title.focus_set()
                return
            text = self._compose_text(title_text, content.get("1.0", tk.END))

            def created(entry: dict[str, Any]) -> None:
                dialog.destroy()
                entry_id = entry.get("id") if isinstance(entry, dict) else None
                self.refresh(select_id=entry_id)

            self._run("新建记录", lambda: self.client.create_entry(text, list(files)), created)

        ttk.Button(bar, text="创建记录", style="Primary.TButton", command=submit).grid(row=0, column=2)
        title.focus_set()

    def _save_text(self) -> None:
        if not self.current or not self.client:
            return
        title = self.title_var.get().strip()
        if not title:
            messagebox.showwarning("缺少标题", "实验标题不能为空。", parent=self)
            self.title_entry.focus_set()
            return
        text = self._compose_current_text()
        entry_id = self.current["id"]

        def saved(result: Any) -> None:
            if isinstance(result, dict) and "id" in result:
                self._show_detail(result)
            self.refresh(select_id=entry_id)

        self._run("保存记录", lambda: self.client.update_entry(entry_id, text), saved)

    def _delete_entry(self) -> None:
        if not self.current or not self.client:
            return
        entry_id = self.current["id"]
        title = self.title_var.get().strip() or f"记录 #{entry_id}"
        if not messagebox.askyesno("删除记录", f"确定删除“{title}”及其全部附件吗？\n\n此操作无法在客户端撤销。", icon="warning", parent=self):
            return

        def deleted(_result: Any) -> None:
            self.current = None
            self._show_empty_state()
            self.refresh()

        self._run("删除记录", lambda: self.client.delete_entry(entry_id), deleted)

    def _selected_attachment(self) -> Optional[dict[str, Any]]:
        if not self.current:
            return None
        selection = self.attachments.selection()
        if not selection:
            return None
        att_id = int(selection[0])
        for item in self.current.get("attachments", []):
            if item.get("id") == att_id:
                return item
        return None

    def _add_attachments(self) -> None:
        if not self.current or not self.client:
            return
        paths = list(filedialog.askopenfilenames(parent=self, title="添加实验附件"))
        if not paths:
            return
        entry_id = self.current["id"]

        def added(result: Any) -> None:
            if isinstance(result, dict) and "id" in result:
                self._show_detail(result)
            else:
                self._run("刷新记录", lambda: self.client.get_entry(entry_id), self._show_detail)
            self.refresh(select_id=entry_id)

        self._run("上传附件", lambda: self.client.add_attachments(entry_id, paths), added)

    def _preview(self) -> None:
        item = self._selected_attachment()
        if not item:
            return
        filename = item.get("filename", "附件")
        if item.get("has_thumb"):
            self._fetch_thumb(item["id"])
        elif filename.lower().endswith(IMAGE_EXTS):
            self.preview_hint_var.set("该图片暂无缩略图\n可下载后查看原图")
            self.preview_label.configure(image="")
            self._photo = None
        else:
            self.preview_hint_var.set(f"{self._file_type(filename)}文件\n{filename}\n\n可点击“下载”查看")
            self.preview_label.configure(image="")
            self._photo = None

    def _fetch_thumb(self, att_id: int) -> None:
        if not self.client:
            return

        def show(data: Optional[bytes]) -> None:
            if not data:
                self.preview_hint_var.set("暂无可用预览")
                self.preview_label.configure(image="")
                return
            try:
                image = Image.open(io.BytesIO(data))
                image.thumbnail((self._px(300), self._px(180)), Image.Resampling.LANCZOS)
                self._photo = ImageTk.PhotoImage(image)
                self.preview_label.configure(image=self._photo)
                self.preview_hint_var.set("")
            except Exception as exc:
                self.preview_hint_var.set(f"预览失败\n{exc}")
                self.preview_label.configure(image="")

        self._run("加载预览", lambda: self.client.get_thumb(att_id), show)

    def _download(self) -> None:
        item = self._selected_attachment()
        if not item or not self.client:
            messagebox.showinfo("选择附件", "请先选择要下载的附件。", parent=self)
            return
        path = filedialog.asksaveasfilename(parent=self, initialfile=item.get("filename", "attachment"))
        if not path:
            return

        def download() -> str:
            data = self.client.download_attachment(item["id"])
            with open(path, "wb") as file:
                file.write(data)
            return path

        self._run("下载附件", download, lambda saved_path: messagebox.showinfo("下载完成", f"文件已保存到：\n{saved_path}", parent=self))

    def _delete_attachment(self) -> None:
        item = self._selected_attachment()
        if not item or not self.client or not self.current:
            return
        if not messagebox.askyesno("删除附件", f"确定删除附件“{item.get('filename', '')}”吗？", icon="warning", parent=self):
            return
        entry_id = self.current["id"]

        def deleted(_result: Any) -> None:
            self._run("刷新记录", lambda: self.client.get_entry(entry_id), self._show_detail)
            self.refresh(select_id=entry_id)

        self._run("删除附件", lambda: self.client.delete_attachment(item["id"]), deleted)

    @staticmethod
    def _file_type(filename: str) -> str:
        ext = os.path.splitext(filename)[1].lower()
        if ext in IMAGE_EXTS:
            return "图片"
        if ext in (".xlsx", ".xls", ".csv", ".tsv"):
            return "数据表"
        if ext in (".pdf", ".doc", ".docx", ".txt", ".md"):
            return "文档"
        return ext[1:].upper() if ext else "文件"

    @staticmethod
    def _format_size(size: Any) -> str:
        try:
            value = max(0, int(size))
        except (TypeError, ValueError):
            value = 0
        if value < 1024:
            return f"{value} B"
        if value < 1024 * 1024:
            return f"{value / 1024:.1f} KB"
        return f"{value / 1024 / 1024:.1f} MB"

    def _on_close(self) -> None:
        if self._dirty and not messagebox.askyesno(
            "尚未保存",
            "当前记录还有未保存的修改，确定直接退出吗？",
            icon="warning",
            parent=self,
        ):
            return
        self.destroy()

    def destroy(self) -> None:
        self._closing = True
        if self.client:
            try:
                self.client.close()
            except Exception:
                pass
        super().destroy()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="ChemDB desktop client")
    parser.add_argument("--demo", action="store_true", help="run with local sample data and no network")
    parser.add_argument("--font-scale", type=float, default=None, help="override the remembered UI font scale (0.9 to 1.4)")
    parser.add_argument("--list-relations", action="store_true", help="print the supported relationship taxonomy and exit")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    if args.list_relations:
        print(json.dumps(RELATION_TAXONOMY, ensure_ascii=False, indent=2))
        raise SystemExit(0)
    ChemDBApp(demo=args.demo, font_scale=args.font_scale).mainloop()
