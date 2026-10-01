import importlib.util
import io
import sys
import unittest
from pathlib import Path


MODULE_PATH = Path(__file__).with_name("main.py")
SPEC = importlib.util.spec_from_file_location("chemdb_main_for_tests", MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


class CoreTests(unittest.TestCase):
    def test_relationship_chain(self):
        entries = [
            {
                "id": 14,
                "text": "关系测试\n物质A与物质B发生反应，物质B与物质C发生吸附，物质A与物质C发生络合。",
            }
        ]
        relations = MODULE.ChemDBApp._extract_relations_from_entries(entries)
        self.assertEqual(
            [(item.source, item.target, item.label) for item in relations],
            [
                ("物质A", "物质B", "反应"),
                ("物质B", "物质C", "吸附"),
                ("物质A", "物质C", "络合"),
            ],
        )

    def test_relationship_in_title_and_body(self):
        entries = [
            {"id": 1, "text": "亚甲基蓝与铅离子吸附实验\n记录不同 pH 条件。"},
            {"id": 2, "text": "正文识别\n刚果红对催化剂的光催化。"},
        ]
        relations = MODULE.ChemDBApp._extract_relations_from_entries(entries)
        self.assertEqual(len(relations), 2)
        self.assertEqual(relations[0].label, "吸附")
        self.assertEqual(relations[1].label, "光催化")

    def test_relationship_terms_from_actual_ui_report(self):
        entries = [
            {
                "id": 17,
                "text": "测试思维导图\n物质A与物质B置换，物质C与物质D螯合，物质X与物质Y嵌合。",
            }
        ]
        relations = MODULE.ChemDBApp._extract_relations_from_entries(entries)
        self.assertEqual(
            [(item.source, item.target, item.label) for item in relations],
            [
                ("物质A", "物质B", "置换"),
                ("物质C", "物质D", "螯合"),
                ("物质X", "物质Y", "嵌合"),
            ],
        )

    def test_relationship_spacing_and_passive_sentence(self):
        entries = [
            {"id": 1, "text": "空格句式\n物质 A 与物质 B 发生置换反应。"},
            {"id": 2, "text": "被动句式\n亚甲基蓝被活性炭吸附。"},
        ]
        relations = MODULE.ChemDBApp._extract_relations_from_entries(entries)
        self.assertEqual(
            [(item.source, item.target, item.label) for item in relations],
            [
                ("物质 A", "物质 B", "置换"),
                ("活性炭", "亚甲基蓝", "吸附"),
            ],
        )

    def test_taxonomy_has_university_chemistry_and_chemical_engineering_coverage(self):
        self.assertGreaterEqual(len(MODULE.RELATION_CATEGORY_BY_LABEL), 120)
        self.assertGreaterEqual(len(MODULE.RELATION_LABEL_ALIASES), 200)
        for category in (
            "基础反应",
            "有机反应",
            "催化与降解",
            "材料与界面",
            "化工分离",
            "传递与流动",
            "影响关系",
            "实验关系",
        ):
            self.assertIn(category, MODULE.RELATION_TAXONOMY)

    def test_direct_catalysis_process_and_interaction_sentences(self):
        entries = [
            {"id": 1, "text": "吸附\n活性炭吸附亚甲基蓝。"},
            {"id": 2, "text": "催化\n二氧化钛光催化亚甲基蓝降解。"},
            {"id": 3, "text": "膜分离\n膜A对染料B进行截留。"},
            {"id": 4, "text": "界面作用\n石墨烯与染料B发生π-π堆积。"},
        ]
        relations = MODULE.ChemDBApp._extract_relations_from_entries(entries)
        self.assertEqual(
            [(item.source, item.target, item.label) for item in relations],
            [
                ("活性炭", "亚甲基蓝", "吸附"),
                ("二氧化钛", "亚甲基蓝", "光催化降解"),
                ("膜A", "染料B", "截留"),
                ("石墨烯", "染料B", "π-π堆积"),
            ],
        )

    def test_condition_and_experiment_relationships(self):
        entries = [
            {"id": 1, "text": "条件趋势\n随着pH升高，去除率提高。"},
            {"id": 2, "text": "条件作用\n升高温度促进硝酸钾溶解。"},
            {"id": 3, "text": "实验优化\n实验B基于实验A进行优化。"},
            {"id": 4, "text": "对照关系\n实验C是实验A的对照实验。"},
            {"id": 5, "text": "复现关系\n实验D复现实验B。"},
        ]
        relations = MODULE.ChemDBApp._extract_relations_from_entries(entries)
        self.assertEqual(
            [(item.source, item.target, item.label) for item in relations],
            [
                ("pH升高", "去除率", "提高"),
                ("升高温度", "硝酸钾", "促进溶解"),
                ("实验B", "实验A", "优化"),
                ("实验C", "实验A", "对照"),
                ("实验D", "实验B", "复现"),
            ],
        )

    def test_graph_spacing_zoom_keeps_mouse_anchor_fixed(self):
        center_x, center_y = 300.0, 240.0
        base_x, base_y = 410.0, 160.0
        old_scale = 1.2
        old_pan_x, old_pan_y = 24.0, -18.0
        anchor_x = center_x + (base_x - center_x) * old_scale + old_pan_x
        anchor_y = center_y + (base_y - center_y) * old_scale + old_pan_y
        new_scale = 1.8
        new_pan_x, new_pan_y = MODULE.ChemDBApp._zoom_pan_for_anchor(
            old_scale,
            new_scale,
            old_pan_x,
            old_pan_y,
            anchor_x,
            anchor_y,
            center_x,
            center_y,
        )
        transformed_x = center_x + (base_x - center_x) * new_scale + new_pan_x
        transformed_y = center_y + (base_y - center_y) * new_scale + new_pan_y
        self.assertAlmostEqual(transformed_x, anchor_x)
        self.assertAlmostEqual(transformed_y, anchor_y)

    def test_graph_spacing_only_changes_coordinates(self):
        app = object.__new__(MODULE.ChemDBApp)
        app._graph_spacing = 2.0
        app._graph_pan_x = 10.0
        app._graph_pan_y = -5.0
        positions = {"A": (250.0, 200.0), "B": (350.0, 200.0)}
        transformed = app._apply_graph_view(positions, 600.0, 400.0)
        self.assertEqual(transformed["A"], (210.0, 195.0))
        self.assertEqual(transformed["B"], (410.0, 195.0))
        self.assertEqual(transformed["B"][0] - transformed["A"][0], 200.0)

    def test_manual_node_offsets_only_move_selected_nodes(self):
        positions = {"Fe³⁺": (300.0, 200.0), "Cu²⁺": (500.0, 200.0)}
        transformed = MODULE.ChemDBApp._apply_node_offsets(positions, {"Fe³⁺": (45.0, -20.0)})
        self.assertEqual(transformed["Fe³⁺"], (345.0, 180.0))
        self.assertEqual(transformed["Cu²⁺"], (500.0, 200.0))

    def test_rich_payload_round_trip(self):
        payload = {
            "version": 1,
            "styles": [
                {
                    "style": {
                        "face": "微软雅黑",
                        "size": 18,
                        "bold": True,
                        "underline": True,
                        "foreground": "#123456",
                        "highlight": "#FFF2A8",
                    },
                    "ranges": [[0, 4]],
                }
            ],
        }
        encoded = "实验正文" + MODULE.RICH_MARKER + MODULE.ChemDBApp._encode_payload(payload) + MODULE.RICH_END
        body, decoded = MODULE.ChemDBApp._decode_rich_body(encoded)
        self.assertEqual(body, "实验正文")
        self.assertEqual(decoded, payload)
        self.assertEqual(MODULE.ChemDBApp._split_text("标题\n" + encoded), ("标题", "实验正文"))

    def test_plain_text_compatibility(self):
        text = "旧版标题\n旧版正文保持不变。"
        self.assertEqual(MODULE.ChemDBApp._split_text(text), ("旧版标题", "旧版正文保持不变。"))

    def test_chemical_notation_and_rich_baselines(self):
        self.assertEqual(MODULE.ChemDBApp._normalize_chemical_notation("Fe^{3+}"), "Fe³⁺")
        self.assertEqual(MODULE.ChemDBApp._normalize_chemical_notation("SO_{4}^{2-}"), "SO₄²⁻")
        payload = {
            "version": 1,
            "styles": [
                {"style": {"baseline": "sub"}, "ranges": [[2, 3]]},
                {"style": {"baseline": "super"}, "ranges": [[3, 5]]},
            ],
        }
        self.assertEqual(MODULE.ChemDBApp._semantic_rich_text("SO42-", payload), "SO₄²⁻")

    def test_custom_keywords_and_importance(self):
        rules = [
            MODULE.KeywordRule("三价铁离子", "concept", "Fe³⁺", 3),
            MODULE.KeywordRule("絮凝", "relation", "絮凝作用", 2),
        ]
        entries = [{"id": 7, "text": "自定义规则\n三价铁离子与OH⁻发生絮凝。"}]
        relations = MODULE.ChemDBApp._extract_relations_from_entries(entries, rules)
        self.assertEqual(
            [(item.source, item.target, item.label, item.importance) for item in relations],
            [("Fe³⁺", "OH⁻", "絮凝作用", 2)],
        )
        priorities = MODULE.ChemDBApp._concept_priority_map(rules)
        self.assertEqual(priorities["fe³⁺"], 3)
        self.assertNotIn("oh⁻", priorities)

    def test_multiple_core_nodes_are_kept_in_central_layer(self):
        core = ["Fe³⁺", "Cu²⁺", "Pb²⁺"]
        outer = ["OH⁻", "活性炭", "亚甲基蓝"]
        positions = MODULE.ChemDBApp._layout_priority_nodes(core, outer, 800, 600, 42)
        self.assertEqual(set(positions), set(core + outer))
        center = (400.0, 300.0)

        def distance(node):
            x, y = positions[node]
            return ((x - center[0]) ** 2 + (y - center[1]) ** 2) ** 0.5

        self.assertLess(max(distance(node) for node in core), min(distance(node) for node in outer))

    def test_csv_attachment_text_extraction(self):
        data = "物质A,物质B,关系\nFe³⁺,OH⁻,沉淀\n".encode("utf-8-sig")
        text = MODULE.ChemDBApp._extract_table_text("result.csv", data)
        self.assertIn("Fe³⁺ OH⁻ 沉淀", text)
        self.assertIn("Fe³⁺与OH⁻发生沉淀", text)
        raw_text = MODULE.ChemDBApp._extract_attachment_text("result.csv", data, for_analysis=False)
        self.assertIn("Fe³⁺ OH⁻ 沉淀", raw_text)
        self.assertNotIn("Fe³⁺与OH⁻发生沉淀", raw_text)
        relations = MODULE.ChemDBApp._extract_relations_from_entries([{"id": 9, "text": "附件\n" + text}])
        self.assertIn(("Fe³⁺", "OH⁻", "沉淀"), [(item.source, item.target, item.label) for item in relations])

    def test_xlsx_attachment_text_extraction(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl unavailable")
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["物质A", "物质B", "关系"])
        sheet.append(["Cu²⁺", "OH⁻", "沉淀"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        text = MODULE.ChemDBApp._extract_table_text("result.xlsx", buffer.getvalue())
        self.assertIn("Cu²⁺ OH⁻ 沉淀", text)

    def test_plain_text_attachment_extraction_and_merge(self):
        extracted = MODULE.ChemDBApp._extract_attachment_text(
            "notes.txt", "Fe³⁺与OH⁻发生沉淀。".encode("utf-8")
        )
        self.assertEqual(extracted, "Fe³⁺与OH⁻发生沉淀。")
        merged = MODULE.ChemDBApp._merge_attachment_text(
            "原有实验记录。", [("notes.txt", extracted)]
        )
        self.assertEqual(
            merged,
            "原有实验记录。\n\n【附件文字：notes.txt】\nFe³⁺与OH⁻发生沉淀。",
        )

    def test_docx_attachment_text_extraction(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx unavailable")
        document = Document()
        document.add_paragraph("Cu²⁺与OH⁻发生沉淀。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "温度"
        table.cell(0, 1).text = "25 ℃"
        buffer = io.BytesIO()
        document.save(buffer)
        text = MODULE.ChemDBApp._extract_attachment_text("notes.docx", buffer.getvalue())
        self.assertIn("Cu²⁺与OH⁻发生沉淀。", text)
        self.assertIn("温度 25 ℃", text)

    def test_attachment_text_merge_with_empty_body_and_multiple_files(self):
        merged = MODULE.ChemDBApp._merge_attachment_text(
            "", [("a.txt", "第一段"), ("b.csv", "第二段")]
        )
        self.assertEqual(
            merged,
            "【附件文字：a.txt】\n第一段\n\n【附件文字：b.csv】\n第二段",
        )


if __name__ == "__main__":
    unittest.main()
