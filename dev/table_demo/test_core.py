import tempfile
import unittest
from pathlib import Path

from core import Sheet, export_tables, load_project, load_tables, make_labels, save_project


class TableTests(unittest.TestCase):
    def test_export(self):
        from openpyxl import load_workbook
        sheet = Sheet('实验', [['A,B', '=1+1']])
        sheet.annotate('row', 1, 1, 2, 'fixed', text='样品')
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)
            export_tables(path / 'out.csv', [sheet])
            self.assertEqual(load_tables(path / 'out.csv')[0].rows, [['A,B 【样品】', '=1+1 【样品】']])
            export_tables(path / 'out.xlsx', [sheet, Sheet('第二张', [['=2+2']])])
            workbook = load_workbook(path / 'out.xlsx')
            self.assertEqual(workbook.sheetnames, ['实验', '第二张'])
            self.assertEqual(workbook['实验']['B1'].value, '=1+1 【样品】')
            self.assertEqual(workbook['第二张']['A1'].data_type, 's')
            workbook.close()
        self.assertEqual(sheet.rows, [['A,B', '=1+1']])

    def test_endpoints(self):
        self.assertEqual(make_labels(3, 'arithmetic', first='1', last='9'), ['1', '5', '9'])
        self.assertEqual(make_labels(3, 'geometric', first='1', last='9'), ['1', '3', '9'])
        self.assertEqual(make_labels(3, 'geometric', first='-9', last='-1'), ['-9', '-3', '-1'])
        self.assertEqual(make_labels(1, 'arithmetic', first='2', last='2'), ['2'])
        with self.assertRaises(ValueError):
            make_labels(1, 'arithmetic', first='1', last='2')
        with self.assertRaises(ValueError):
            make_labels(3, 'geometric', first='0', last='9')

    def test_sequences(self):
        self.assertEqual(make_labels(3, 'arithmetic', '浓度=', '0.1', '0.1', ' mM'), ['浓度=0.1 mM', '浓度=0.2 mM', '浓度=0.3 mM'])
        self.assertEqual(make_labels(4, 'geometric', '', '8', '0.5'), ['8', '4', '2', '1'])
        self.assertEqual(make_labels(3, 'arithmetic', '', '1', '-1'), ['1', '0', '-1'])
        self.assertEqual(make_labels(3, 'geometric', '', '2', '0'), ['2', '0', '0'])
        with self.assertRaises(ValueError):
            make_labels(2, 'arithmetic', first='NaN')

    def test_annotations_and_roundtrip(self):
        sheet = Sheet('实验', [['a', 'b'], ['c', 'd'], ['e', 'f']])
        sheet.annotate('row', 2, 1, 2, 'fixed', text='实验组')
        sheet.annotate('column', 2, 2, 3, 'arithmetic', first='0.1', step='0.1')
        sheet.annotate('row', 2, 1, 2, 'fixed', text='实验组')
        self.assertEqual(sheet.labels[2, 2], ['实验组', '0.1'])
        self.assertEqual(sheet.rows[1], ['c', 'd'])
        with self.assertRaises(ValueError):
            sheet.annotate('row', 0, 1, 2, 'fixed', text='bad')
        self.assertNotIn((0, 1), sheet.labels)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'project.json'
            save_project(path, [sheet])
            self.assertEqual(load_project(path), [sheet])

    def test_csv_xml_and_xlsx(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)
            (path / 'a.csv').write_text('名称,值\n"A,B",0.1\nC\n', encoding='utf-8-sig')
            self.assertEqual(load_tables(path / 'a.csv')[0].rows, [['名称', '值'], ['A,B', '0.1'], ['C', '']])
            (path / 'a.tsv').write_text('名称\t值\nA\t2', encoding='gb18030')
            self.assertEqual(load_tables(path / 'a.tsv')[0].rows[1], ['A', '2'])
            (path / 'a.xml').write_text('<records><record><name>A</name><value>1</value></record><record><name>B</name></record></records>', encoding='utf-8')
            self.assertEqual(load_tables(path / 'a.xml')[0].rows, [['name', 'value'], ['A', '1'], ['B', '']])
            (path / 'a.xml').write_text('<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="Test"><Table><Row ss:Index="2"><Cell ss:Index="3"><Data ss:Type="String">X</Data></Cell></Row></Table></Worksheet></Workbook>', encoding='utf-8')
            self.assertEqual(load_tables(path / 'a.xml')[0].rows, [['', '', ''], ['', '', 'X']])
            workbook = Workbook()
            workbook.active.append(['样品', 1])
            workbook.create_sheet('第二张').append(['=1+1'])
            workbook.save(path / 'a.xlsx')
            sheets = load_tables(path / 'a.xlsx')
            self.assertEqual(len(sheets), 2)
            self.assertEqual(sheets[1].rows, [['=1+1']])

    def test_xml_entity_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'bad.xml'
            path.write_text('<!DOCTYPE x [<!ENTITY a "text">]><table><row><cell>&a;</cell></row></table>')
            from defusedxml.common import DefusedXmlException
            with self.assertRaises(DefusedXmlException):
                load_tables(path)


if __name__ == '__main__':
    unittest.main()
