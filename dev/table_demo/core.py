"""UI-independent table import and cell annotations; coordinates are 1-based."""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path


@dataclass
class Sheet:
    name: str
    rows: list[list[str]]
    labels: dict[tuple[int, int], list[str]] = field(default_factory=dict)

    def __post_init__(self):
        width = max((len(row) for row in self.rows), default=0)
        self.rows = [row + [''] * (width - len(row)) for row in self.rows]

    @property
    def width(self):
        return len(self.rows[0]) if self.rows else 0

    def annotate(self, axis, index, start, end, mode, text='', first='1', step='1', suffix='', last=None):
        if axis not in ('row', 'column'):
            raise ValueError('请选择行或列')
        limit, length = (len(self.rows), self.width) if axis == 'row' else (self.width, len(self.rows))
        if not 1 <= index <= limit or not 1 <= start <= end <= length:
            raise ValueError('行列编号或范围超出表格边界（编号从 1 开始）')
        tags = make_labels(end - start + 1, mode, text, first, step, suffix, last=last)
        for position, tag in zip(range(start, end + 1), tags):
            cell = (index, position) if axis == 'row' else (position, index)
            existing = self.labels.setdefault(cell, [])
            if tag not in existing:
                existing.append(tag)


def make_labels(count, mode, text='', first='1', step='1', suffix='', last=None):
    if mode == 'fixed':
        if not text.strip():
            raise ValueError('固定标签不能为空')
        return [text] * count
    if mode not in ('arithmetic', 'geometric'):
        raise ValueError('未知的标签模式')
    try:
        value, delta = Decimal(first), Decimal(step)
        if not value.is_finite() or not delta.is_finite():
            raise ValueError('首项和公差/公比必须是有限数字')
        result = []
        with localcontext() as context:
            context.prec = 50
            endpoint = None
            if last is not None:
                endpoint = Decimal(last)
                if not endpoint.is_finite():
                    raise ValueError('终点必须是有限数字')
                if count == 1:
                    if value != endpoint:
                        raise ValueError('只有一个格子时，起点和终点必须相同')
                elif mode == 'arithmetic':
                    delta = (endpoint - value) / (count - 1)
                else:
                    if value == 0 or endpoint == 0 or (value > 0) != (endpoint > 0):
                        raise ValueError('等比起止模式要求起点和终点非零且同号；零或交替正负请使用首项＋公比')
                    delta = (endpoint / value) ** (Decimal(1) / (count - 1))
            for i in range(count):
                if endpoint is not None and i == count - 1:
                    value = endpoint
                if not value.is_finite() or (value and abs(value.adjusted()) > 1000):
                    raise ValueError('数列数值过大或过小')
                number = format(value, 'f')
                if '.' in number:
                    number = number.rstrip('0').rstrip('.')
                result.append(text + ('0' if value == 0 else number) + suffix)
                if i + 1 < count:
                    value = value + delta if mode == 'arithmetic' else value * delta
        return result
    except (InvalidOperation, ArithmeticError) as exc:
        raise ValueError('请输入有效的首项和公差/公比') from exc


def _text(value):
    return '' if value is None else str(value)


def load_tables(filename):
    path = Path(filename)
    ext = path.suffix.lower()
    if ext in ('.csv', '.tsv', '.txt'):
        raw = path.read_bytes()
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = raw.decode('gb18030')
        delimiter = '\t' if ext == '.tsv' else ','
        try:
            delimiter = csv.Sniffer().sniff(content[:8192], delimiters=',\t;|').delimiter
        except csv.Error:
            pass
        import io
        sheets = [Sheet(path.stem, list(csv.reader(io.StringIO(content), delimiter=delimiter)))]
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheets = [Sheet(ws.title, [[_text(v) for v in row] for row in ws.iter_rows(values_only=True)]) for ws in workbook]
        finally:
            workbook.close()
    elif ext == '.xls':
        import xlrd
        workbook = xlrd.open_workbook(str(path))
        try:
            sheets = [Sheet(ws.name, [[_text(v) for v in ws.row_values(r)] for r in range(ws.nrows)]) for ws in workbook.sheets()]
        finally:
            workbook.release_resources()
    elif ext == '.xml':
        sheets = _load_xml(path)
    else:
        raise ValueError('支持 xlsx、xls、csv、tsv、txt、xml')
    sheets = [sheet for sheet in sheets if sheet.rows and sheet.width]
    if not sheets:
        raise ValueError('文件中没有可读取的表格')
    return sheets


def _load_xml(path):
    from defusedxml import ElementTree as ET
    root = ET.parse(path).getroot()
    ns = '{urn:schemas-microsoft-com:office:spreadsheet}'
    if root.tag == ns + 'Workbook':
        sheets = []
        for ws in root.findall(ns + 'Worksheet'):
            rows = []
            for row in ws.findall('./' + ns + 'Table/' + ns + 'Row'):
                row_index = int(row.get(ns + 'Index', len(rows) + 1))
                if not len(rows) < row_index <= 100000:
                    raise ValueError('XML 行索引无效或超过 100000')
                rows.extend([[] for _ in range(row_index - len(rows) - 1)])
                values = []
                for cell in row.findall(ns + 'Cell'):
                    index = int(cell.get(ns + 'Index', len(values) + 1))
                    if not len(values) < index <= 16384:
                        raise ValueError('XML 列索引无效或超过 16384')
                    values.extend([''] * (index - len(values) - 1))
                    data = cell.find(ns + 'Data')
                    values.append(''.join(data.itertext()) if data is not None else '')
                rows.append(values)
            sheets.append(Sheet(ws.get(ns + 'Name', 'Sheet'), rows))
        return sheets
    records = list(root)
    if not records or any(not list(record) for record in records):
        raise ValueError('XML 需为 Excel 2003 XML，或根元素下每条记录包含字段元素的结构')
    if any(list(cell) for record in records for cell in record):
        raise ValueError('暂不支持嵌套字段 XML，请转换为平面记录')
    if all(record.tag.rsplit('}', 1)[-1].lower() == 'row' for record in records) and all(cell.tag.rsplit('}', 1)[-1].lower() == 'cell' for record in records for cell in record):
        return [Sheet(path.stem, [[cell.text or '' for cell in row] for row in records])]
    headers = list(dict.fromkeys(cell.tag for record in records for cell in record))
    rows = [[h.rsplit('}', 1)[-1] for h in headers]]
    for record in records:
        fields = {cell.tag: cell.text or '' for cell in record}
        if len(fields) != len(record):
            raise ValueError('同一记录有重复字段，无法自动映射为列')
        rows.append([fields.get(h, '') for h in headers])
    return [Sheet(path.stem, rows)]


def export_tables(filename, sheets):
    """Export displayed values as text; XLSX includes all supplied sheets."""
    path = Path(filename)
    def rows(sheet):
        for r, row in enumerate(sheet.rows, 1):
            yield [value + ''.join(f' 【{tag}】' for tag in sheet.labels.get((r, c), []))
                   for c, value in enumerate(row, 1)]

    if path.suffix.lower() == '.csv':
        if len(sheets) != 1:
            raise ValueError('CSV 一次只能导出一个工作表')
        with path.open('w', encoding='utf-8-sig', newline='') as output:
            csv.writer(output).writerows(rows(sheets[0]))
    elif path.suffix.lower() == '.xlsx':
        import re
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet in sheets:
            title = re.sub(r'[\\/*?:\[\]]', '_', sheet.name)[:31] or 'Sheet'
            ws = workbook.create_sheet(title)
            for r, row in enumerate(rows(sheet), 1):
                for c, value in enumerate(row, 1):
                    if len(value) > 32767:
                        raise ValueError(f'{sheet.name} 第 {r} 行第 {c} 列超出 Excel 单元格文字长度限制')
                    cell = ws.cell(r, c, value)
                    cell.data_type = 's'
        workbook.save(path)
        workbook.close()
    else:
        raise ValueError('请选择 .xlsx 或 .csv 格式')


def save_project(filename, sheets):
    payload = {'version': 1, 'sheets': [
        {'name': s.name, 'rows': s.rows, 'labels': [
            {'row': r, 'column': c, 'tags': tags} for (r, c), tags in sorted(s.labels.items())]}
        for s in sheets]}
    Path(filename).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')


def load_project(filename):
    payload = json.loads(Path(filename).read_text(encoding='utf-8'))
    if payload['version'] != 1:
        raise ValueError('不支持的项目版本')
    sheets = []
    for item in payload['sheets']:
        if not isinstance(item['rows'], list) or any(not isinstance(row, list) or any(not isinstance(v, str) for v in row) for row in item['rows']):
            raise ValueError('项目表格格式错误')
        sheet = Sheet(item['name'], item['rows'])
        for label in item['labels']:
            r, c, tags = label['row'], label['column'], label['tags']
            if not (1 <= r <= len(sheet.rows) and 1 <= c <= sheet.width) or not isinstance(tags, list) or any(not isinstance(t, str) for t in tags):
                raise ValueError('项目标签格式错误')
            sheet.labels[r, c] = tags
        sheets.append(sheet)
    if not sheets:
        raise ValueError('项目中没有工作表')
    return sheets
